import os
import re
import json
import base64
import io
import urllib.parse
import docx
import difflib
import docx.shared
from http.server import ThreadingHTTPServer, HTTPServer, BaseHTTPRequestHandler
from fpdf import FPDF
from PIL import Image

try:
    import curriculum_service
except ImportError:
    import sys
    sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
    import curriculum_service

try:
    from qrcodegen import QrCode
except ImportError:
    import sys
    sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
    from qrcodegen import QrCode

def clean_pdf_text(text):
    if not text:
        return ""
    text = str(text)
    replacements = {
        "\u2014": "--",
        "\u2013": "-",
        "\u2018": "'",
        "\u2019": "'",
        "\u201c": '"',
        "\u201d": '"',
        "\u2022": "*",
        "\u2026": "...",
        "\u00a0": " ",
        "–": "-",
        "—": "--",
        "“": '"',
        "”": '"',
        "‘": "'",
        "’": "'",
        "\u2713": "[OK]",
        "\u2714": "[OK]",
        "\u2717": "[X]",
        "\u2718": "[X]",
    }
    for orig, repl in replacements.items():
        text = text.replace(orig, repl)
    return text.encode("latin-1", "replace").decode("latin-1")

def wrap_text(pdf, text, max_w):
    words = clean_pdf_text(text).split()
    if not words:
        return [""]
    lines, cur = [], []
    for w in words:
        test = " ".join(cur + [w])
        if pdf.get_string_width(test) <= max_w:
            cur.append(w)
        else:
            if cur:
                lines.append(" ".join(cur))
                cur = [w]
            else:
                lines.append(w)
                cur = []
    if cur:
        lines.append(" ".join(cur))
    return lines

def generate_qr_png_base64(url):
    qr = QrCode.encode_text(url, QrCode.Ecc.MEDIUM)
    border = 4
    scale = 8
    n = qr.get_size()
    dim = (n + border * 2) * scale
    img = Image.new('RGB', (dim, dim), 'white')
    pix = img.load()
    for y in range(n):
        for x in range(n):
            if qr.get_module(x, y):
                for dy in range(scale):
                    for dx in range(scale):
                        pix[(x + border) * scale + dx, (y + border) * scale + dy] = (0, 9, 83)
    buf = io.BytesIO()
    img.save(buf, format='PNG')
    return 'data:image/png;base64,' + base64.b64encode(buf.getvalue()).decode('utf-8')

def generate_door_poster(payload):
    pdf = FPDF(orientation='P', unit='mm', format='A4')
    pdf.set_margins(12, 10, 12)
    pdf.set_auto_page_break(False)
    pdf.add_page()
    
    # Outer double border (MTTI navy & gold)
    pdf.set_draw_color(0, 9, 83)
    pdf.set_line_width(0.8)
    pdf.rect(8, 8, 194, 281)
    pdf.set_draw_color(196, 136, 32)
    pdf.set_line_width(0.4)
    pdf.rect(9.5, 9.5, 191, 278)
    
    # Header
    pdf.set_font('Helvetica', 'B', 15)
    pdf.set_text_color(0, 9, 83)
    pdf.set_xy(12, 14)
    pdf.cell(186, 7, 'MUKIRIA TECHNICAL TRAINING INSTITUTE', ln=1, align='C')
    
    dept = clean_pdf_text(payload.get('department', 'DEPARTMENT OF COMPUTING AND INFORMATICS').upper())
    pdf.set_font('Helvetica', 'B', 10)
    pdf.set_text_color(80, 80, 80)
    pdf.cell(186, 5, dept, ln=1, align='C')
    
    pdf.set_font('Helvetica', '', 8)
    pdf.set_text_color(120, 120, 120)
    pdf.cell(186, 4, 'ISO 9001:2015 CERTIFIED | TVET CDACC CURRICULUM DELIVERY', ln=1, align='C')
    
    # Title Banner
    pdf.ln(2)
    pdf.set_fill_color(0, 9, 83)
    pdf.set_text_color(255, 255, 255)
    pdf.set_font('Helvetica', 'B', 13)
    pdf.cell(186, 9, 'WORKSHOP DOOR ACCESS & LEARNING OUTCOMES', ln=1, align='C', fill=True)
    
    # Session Details Box
    pdf.ln(3)
    pdf.set_fill_color(248, 250, 252)
    pdf.set_draw_color(200, 200, 200)
    pdf.set_line_width(0.3)
    box_y = pdf.get_y()
    pdf.rect(12, box_y, 186, 36, 'DF')
    
    unit_code = clean_pdf_text(payload.get('unit_code', ''))
    unit_name = clean_pdf_text(payload.get('unit_name', ''))
    session_title = clean_pdf_text(payload.get('session_title', 'Practical Laboratory Session'))
    class_code = clean_pdf_text(payload.get('class_code', 'EE6/M/S/24'))
    date_str = clean_pdf_text(payload.get('date', '02/10/2026'))
    time_dur = clean_pdf_text(payload.get('time_duration', '10:30 - 12:30'))
    venue = clean_pdf_text(payload.get('venue', 'Computer Lab 1 / Workshop'))
    trainer = clean_pdf_text(payload.get('trainer_name', 'Alexander Kinoti'))
    
    pdf.set_text_color(0, 9, 83)
    pdf.set_font('Helvetica', 'B', 10)
    pdf.set_xy(15, box_y + 3)
    pdf.cell(180, 5, f'UNIT: {unit_code} - {unit_name.upper()}', ln=1)
    
    pdf.set_font('Helvetica', 'B', 11)
    pdf.set_text_color(196, 136, 32)
    pdf.set_x(15)
    pdf.cell(180, 5.5, f'TOPIC: {session_title}', ln=1)
    
    pdf.set_font('Helvetica', '', 9)
    pdf.set_text_color(40, 40, 40)
    pdf.set_x(15)
    pdf.cell(90, 5, f'Class: {class_code}', ln=0)
    pdf.cell(90, 5, f'Date: {date_str}', ln=1)
    
    pdf.set_x(15)
    pdf.cell(90, 5, f'Time: {time_dur}', ln=0)
    pdf.cell(90, 5, f'Venue: {venue}', ln=1)
    
    pdf.set_x(15)
    pdf.cell(180, 5, f'Trainer: {trainer}', ln=1)
    
    # QR Code Section
    qr_y = box_y + 40
    pdf.set_xy(12, qr_y)
    
    # Generate QR
    target_url = payload.get('target_url', f"http://localhost:3000/session/{payload.get('session_plan_id', 'sp-1')}")
    qr = QrCode.encode_text(target_url, QrCode.Ecc.MEDIUM)
    qr_size_mm = 65
    qr_x = (210 - qr_size_mm) / 2
    
    # Draw QR frame
    pdf.set_fill_color(255, 255, 255)
    pdf.set_draw_color(0, 9, 83)
    pdf.set_line_width(0.5)
    pdf.rect(qr_x - 4, qr_y - 2, qr_size_mm + 8, qr_size_mm + 8, 'DF')
    
    # Vector draw QR
    n = qr.get_size()
    mod_size = qr_size_mm / n
    pdf.set_fill_color(0, 9, 83)
    for r in range(n):
        for c in range(n):
            if qr.get_module(c, r):
                pdf.rect(qr_x + c * mod_size, qr_y + 2 + r * mod_size, mod_size, mod_size, 'F')
                
    # QR Subtext
    sub_y = qr_y + qr_size_mm + 10
    pdf.set_xy(12, sub_y)
    pdf.set_font('Helvetica', 'B', 10)
    pdf.set_text_color(0, 9, 83)
    pdf.cell(186, 5, "SCAN WITH SMARTPHONE CAMERA FOR TODAY'S LEARNING OUTCOMES", ln=1, align='C')
    
    pdf.set_font('Helvetica', '', 8)
    pdf.set_text_color(100, 100, 100)
    pdf.cell(186, 4, clean_pdf_text(f'Direct Link: {target_url}'), ln=1, align='C')
    
    # Learning Outcomes Box
    outcomes_y = sub_y + 11
    pdf.set_xy(12, outcomes_y)
    pdf.set_fill_color(240, 245, 255)
    pdf.set_draw_color(0, 9, 83)
    pdf.set_line_width(0.4)
    
    outcomes = payload.get('learning_outcomes', ['Master practical skills matching CDACC criteria.'])
    box_h = 10 + len(outcomes) * 8
    pdf.rect(12, outcomes_y, 186, box_h, 'DF')
    
    pdf.set_xy(16, outcomes_y + 3)
    pdf.set_font('Helvetica', 'B', 10)
    pdf.set_text_color(0, 9, 83)
    pdf.cell(178, 5, "TODAY'S SPECIFIC LEARNING OUTCOMES (MTTI/F/CUR/05):", ln=1)
    
    pdf.set_font('Helvetica', '', 9.5)
    pdf.set_text_color(20, 20, 20)
    for idx, out in enumerate(outcomes):
        pdf.set_x(18)
        pdf.cell(6, 6, f'{idx+1}.', ln=0)
        pdf.multi_cell(170, 6, clean_pdf_text(str(out)))
        
    # Safety notice at bottom
    safety_y = 265
    pdf.set_xy(12, safety_y)
    pdf.set_fill_color(255, 245, 230)
    pdf.set_draw_color(196, 136, 32)
    pdf.rect(12, safety_y, 186, 18, 'DF')
    
    pdf.set_xy(15, safety_y + 2)
    pdf.set_font('Helvetica', 'B', 8.5)
    pdf.set_text_color(196, 136, 32)
    pdf.cell(180, 4, 'WORKSHOP SAFETY & LABORATORY ETIQUETTE:', ln=1)
    
    pdf.set_font('Helvetica', '', 7.5)
    pdf.set_text_color(60, 60, 60)
    safety_text = clean_pdf_text(payload.get('safety_requirements', 'Wear protective laboratory gear | Strictly no food or beverages in the computer lab | Observe electrical safety rules | Leave workstation orderly.'))
    pdf.set_x(15)
    pdf.multi_cell(180, 3.8, safety_text)
    
    out = pdf.output(dest='S')
    pdf_bytes = out.encode('latin1') if isinstance(out, str) else bytes(out)
    return pdf_bytes

def generate_row_book(payload):
    term = clean_pdf_text(payload.get("term", "Term 3, 2026"))
    trainer = clean_pdf_text(payload.get("trainer_name", "Alexander Kinoti"))
    department = clean_pdf_text(payload.get("department", "Department of Computing and Informatics"))
    records = payload.get("records", [])
    
    pdf = FPDF(orientation='P', unit='mm', format='A4')
    pdf.set_font('Helvetica', '', 10)
    pdf.set_margins(12, 10, 12)
    pdf.set_auto_page_break(False)
    
    current_dir = os.path.dirname(os.path.abspath(__file__))
    project_root = os.path.dirname(current_dir)
    logo_path = os.path.join(project_root, "client", "public", "mtti-logo.jpg")
    
    # Page 1: Cover Page
    pdf.add_page(orientation='P')
    pdf.set_draw_color(0, 9, 83)
    pdf.set_line_width(1.0)
    pdf.rect(10, 10, 190, 277)
    pdf.set_draw_color(196, 136, 32)
    pdf.set_line_width(0.5)
    pdf.rect(12, 12, 186, 273)
    
    y_hdr = 18
    if os.path.exists(logo_path):
        try:
            pdf.image(logo_path, x=92, y=y_hdr, w=26)
            y_hdr += 28
        except Exception:
            y_hdr = 20
    
    pdf.set_xy(15, y_hdr)
    pdf.set_font('Helvetica', 'B', 16)
    pdf.set_text_color(0, 9, 83)
    pdf.cell(180, 7, 'MUKIRIA TECHNICAL TRAINING INSTITUTE', ln=1, align='C')
    
    pdf.set_font('Helvetica', 'B', 11)
    pdf.set_text_color(80, 80, 80)
    pdf.cell(180, 5.5, clean_pdf_text(department.upper()), ln=1, align='C')
    
    pdf.set_font('Helvetica', '', 8.5)
    pdf.set_text_color(120, 120, 120)
    pdf.cell(180, 4.5, 'ISO 9001:2015 CERTIFIED | CURRICULUM IMPLEMENTATION RECORDS', ln=1, align='C')
    
    pdf.ln(5)
    pdf.set_fill_color(0, 9, 83)
    pdf.set_text_color(255, 255, 255)
    pdf.set_font('Helvetica', 'B', 13)
    pdf.cell(180, 10, 'RECORD OF WORK (RoW) COMPILATION BOOK', ln=1, align='C', fill=True)
    
    pdf.set_fill_color(196, 136, 32)
    pdf.set_font('Helvetica', 'B', 9)
    pdf.cell(180, 5.5, clean_pdf_text('OFFICIAL FORM: MTTI/F/CUR/02 -- END OF TERM MASTER DOSSIER'), ln=1, align='C', fill=True)
    
    pdf.ln(5)
    pdf.set_fill_color(248, 250, 252)
    pdf.set_draw_color(200, 200, 200)
    pdf.set_line_width(0.3)
    m_y = pdf.get_y()
    pdf.rect(15, m_y, 180, 36, 'DF')
    
    pdf.set_text_color(0, 9, 83)
    pdf.set_font('Helvetica', 'B', 9.5)
    pdf.set_xy(20, m_y + 3)
    pdf.cell(85, 5.5, clean_pdf_text(f'ACADEMIC TERM: {term}'), ln=0)
    pdf.cell(85, 5.5, clean_pdf_text(f'DATE COMPILED: {payload.get("compiled_date", "October 2026")}'), ln=1)
    
    pdf.set_x(20)
    pdf.cell(85, 5.5, clean_pdf_text(f'TRAINER: {trainer.upper()}'), ln=0)
    pdf.cell(85, 5.5, clean_pdf_text(f'STAFF NO: {payload.get("trainer_id", "MTTI/TR/088")}'), ln=1)
    
    pdf.set_x(20)
    pdf.cell(85, 5.5, 'INSTITUTION: Mukiria TTI (Meru)', ln=0)
    pdf.cell(85, 5.5, 'REGULATORY BODY: TVET CDACC', ln=1)
    
    pdf.set_x(20)
    pdf.cell(170, 5.5, 'COVERAGE PERIOD: Full Academic Term Cycle (Weeks 1 - 12)', ln=1)
    
    tot_sessions = len(records)
    tot_hours = sum(float(r.get("hours_covered", 0) or 0) for r in records)
    unique_units = list(dict.fromkeys(r.get("unit_code", "") for r in records if r.get("unit_code")))
    avg_att = round(sum(float(r.get("trainees_present", 0) or 0) for r in records) / max(1, tot_sessions))
    
    pdf.set_xy(15, m_y + 42)
    pdf.set_font('Helvetica', 'B', 10.5)
    pdf.set_text_color(0, 9, 83)
    pdf.cell(180, 5.5, 'EXECUTIVE TERM METRICS SUMMARY', ln=1)
    
    box_w = 42
    box_h = 18
    box_start_y = pdf.get_y() + 1
    
    stats = [
        ("Delivered Sessions", f"{tot_sessions} Logged", (0, 9, 83)),
        ("Contact Hours", f"{tot_hours:g} Hours", (16, 120, 70)),
        ("Units Covered", f"{len(unique_units)} Units", (196, 136, 32)),
        ("Avg Attendance", f"{avg_att} Trainees", (90, 40, 140))
    ]
    
    for i, (label, val, col) in enumerate(stats):
        bx = 15 + i * (box_w + 4)
        pdf.set_fill_color(245, 247, 250)
        pdf.set_draw_color(col[0], col[1], col[2])
        pdf.set_line_width(0.6)
        pdf.rect(bx, box_start_y, box_w, box_h, 'DF')
        
        pdf.set_xy(bx, box_start_y + 1.5)
        pdf.set_font('Helvetica', '', 7)
        pdf.set_text_color(100, 100, 100)
        pdf.cell(box_w, 3.5, label, ln=1, align='C')
        
        pdf.set_xy(bx, box_start_y + 6.5)
        pdf.set_font('Helvetica', 'B', 10)
        pdf.set_text_color(col[0], col[1], col[2])
        pdf.cell(box_w, 8, val, ln=1, align='C')
        
    pdf.set_y(box_start_y + box_h + 6)
    
    pdf.set_font('Helvetica', 'B', 9.5)
    pdf.set_text_color(0, 9, 83)
    pdf.cell(180, 5, 'SUMMARY BREAKDOWN BY UNIT OF COMPETENCE', ln=1)
    
    pdf.set_fill_color(0, 9, 83)
    pdf.set_text_color(255, 255, 255)
    pdf.set_font('Helvetica', 'B', 7.5)
    t_widths = [12, 45, 75, 25, 23]
    t_headers = ["No.", "Unit Code", "Unit Title / Description", "Sessions", "Delivered Hrs"]
    
    cur_x = 15
    cur_y = pdf.get_y()
    for w, h in zip(t_widths, t_headers):
        pdf.rect(cur_x, cur_y, w, 5.5, 'F')
        pdf.set_xy(cur_x, cur_y + 0.8)
        pdf.cell(w, 4, h, align='C')
        cur_x += w
    pdf.set_y(cur_y + 5.5)
    
    pdf.set_font('Helvetica', '', 7.5)
    pdf.set_text_color(30, 30, 30)
    
    for idx, u_code in enumerate(unique_units):
        u_records = [r for r in records if r.get("unit_code") == u_code]
        u_hours = sum(float(r.get("hours_covered", 0) or 0) for r in u_records)
        u_title = u_records[0].get("unit_name", "") if u_records else ""
        if not u_title:
            u_title = "Competency Standard Unit"
            
        r_y = pdf.get_y()
        r_fill = (idx % 2 == 1)
        pdf.set_fill_color(248, 248, 250) if r_fill else pdf.set_fill_color(255, 255, 255)
        pdf.set_draw_color(220, 220, 220)
        pdf.set_line_width(0.2)
        
        row_vals = [str(idx + 1), clean_pdf_text(u_code), clean_pdf_text(u_title[:45]), str(len(u_records)), f"{u_hours:g} hrs"]
        cur_x = 15
        for w, val, align in zip(t_widths, row_vals, ['C', 'L', 'L', 'C', 'C']):
            pdf.rect(cur_x, r_y, w, 5, 'DF' if r_fill else 'D')
            pdf.set_xy(cur_x + 1, r_y + 0.8)
            pdf.cell(w - 2, 3.5, val, align=align)
            cur_x += w
        pdf.set_y(r_y + 5)
        
    sig_y = 246
    pdf.set_xy(15, sig_y)
    pdf.set_draw_color(0, 9, 83)
    pdf.set_line_width(0.5)
    pdf.line(15, sig_y, 195, sig_y)
    
    pdf.set_xy(15, sig_y + 2)
    pdf.set_font('Helvetica', 'B', 8)
    pdf.set_text_color(0, 9, 83)
    pdf.cell(60, 4, '1. Trainer Verification:', ln=0)
    pdf.cell(60, 4, '2. Head of Department (HOD):', ln=0)
    pdf.cell(60, 4, '3. Quality Assurance / DPA:', ln=1)
    
    pdf.set_font('Helvetica', '', 7)
    pdf.set_text_color(60, 60, 60)
    pdf.cell(60, 3.5, clean_pdf_text(f'Name: {trainer.upper()}'), ln=0)
    pdf.cell(60, 3.5, 'Name: HOD COMPUTING', ln=0)
    pdf.cell(60, 3.5, 'Name: DPA / QUALITY ASSURANCE', ln=1)
    
    pdf.cell(60, 3.5, 'Sign: ____________________', ln=0)
    pdf.cell(60, 3.5, 'Sign: ____________________', ln=0)
    pdf.cell(60, 3.5, 'Sign: ____________________', ln=1)
    
    pdf.cell(60, 3.5, clean_pdf_text(f'Date: {payload.get("compiled_date", "October 2026")}'), ln=0)
    pdf.cell(60, 3.5, 'Date: ____________________', ln=0)
    pdf.cell(60, 3.5, 'Date: ____________________', ln=1)
    
    # Pages 2+: Landscape Records of Work
    col_widths = [14, 22, 34, 106, 14, 18, 47, 22]
    headers = ["Wk", "Date", "Unit & Class", "Work Actually Covered", "Hrs", "Present", "Lesson Reflection / Remarks", "Sign & Date"]
    
    sorted_records = sorted(records, key=lambda r: (str(r.get("unit_code", "")), int(r.get("week_number", 0) or 0)))
    
    def render_landscape_header(pdf_inst, current_unit="", p_num=2):
        pdf_inst.add_page(orientation='L')
        pdf_inst.set_xy(10, 7)
        pdf_inst.set_font('Helvetica', 'B', 12)
        pdf_inst.set_text_color(0, 9, 83)
        pdf_inst.cell(190, 5, 'MUKIRIA TECHNICAL TRAINING INSTITUTE', ln=0)
        
        pdf_inst.set_font('Helvetica', 'B', 8.5)
        pdf_inst.set_text_color(196, 136, 32)
        pdf_inst.cell(87, 5, 'FORM: MTTI/F/CUR/02 | ISO 9001:2015', ln=1, align='R')
        
        pdf_inst.set_font('Helvetica', 'B', 8.5)
        pdf_inst.set_text_color(60, 60, 60)
        unit_str = f" | UNIT: {clean_pdf_text(current_unit)}" if current_unit else ""
        pdf_inst.cell(190, 4, clean_pdf_text(f'RECORD OF WORK (RoW) COMPILATION -- {term.upper()} | TRAINER: {trainer.upper()}{unit_str}'), ln=0)
        pdf_inst.set_font('Helvetica', '', 7.5)
        pdf_inst.cell(87, 4, f'Page {p_num}', ln=1, align='R')
        
        pdf_inst.ln(1.5)
        pdf_inst.set_fill_color(0, 9, 83)
        pdf_inst.set_text_color(255, 255, 255)
        pdf_inst.set_font('Helvetica', 'B', 7.5)
        
        cur_x = 10
        cur_y = pdf_inst.get_y()
        for w, h in zip(col_widths, headers):
            pdf_inst.rect(cur_x, cur_y, w, 6, 'F')
            pdf_inst.set_xy(cur_x, cur_y + 1)
            pdf_inst.cell(w, 4, h, align='C')
            cur_x += w
        pdf_inst.set_xy(10, cur_y + 6)
        
    page_count = 2
    render_landscape_header(pdf, p_num=page_count)
    
    pdf.set_font('Helvetica', '', 7.5)
    pdf.set_text_color(20, 20, 20)
    line_h = 3.6
    
    for row_idx, r in enumerate(sorted_records):
        wk = str(r.get("week_number", ""))
        dt = str(r.get("date_delivered", ""))
        u_info = f"{clean_pdf_text(r.get('unit_code', ''))}\nClass: {clean_pdf_text(r.get('class_code', ''))}"
        work = clean_pdf_text(r.get("work_actually_covered", ""))
        hrs = str(r.get("hours_covered", "2"))
        att = str(r.get("trainees_present", "28"))
        refl = clean_pdf_text(r.get("reflection", "Delivered matching syllabus outcomes."))
        sig_trainer = clean_pdf_text(r.get('signature', trainer.split()[-1]))
        sig = f"{sig_trainer}\n{clean_pdf_text(r.get('signature_date', dt))}"
        
        texts = [wk, dt, u_info, work, hrs, att, refl, sig]
        aligns = ['C', 'C', 'L', 'L', 'C', 'C', 'L', 'C']
        
        all_lines = [wrap_text(pdf, t, w - 2) for t, w in zip(texts, col_widths)]
        max_lines = max(len(l) for l in all_lines)
        row_h = max_lines * line_h + 2.5
        
        if pdf.get_y() + row_h > pdf.h - 18:
            page_count += 1
            render_landscape_header(pdf, p_num=page_count)
            pdf.set_font('Helvetica', '', 7.5)
            pdf.set_text_color(20, 20, 20)
            
        cur_y = pdf.get_y()
        cur_x = 10
        fill = (row_idx % 2 == 1)
        pdf.set_fill_color(250, 250, 252) if fill else pdf.set_fill_color(255, 255, 255)
        pdf.set_draw_color(200, 200, 200)
        pdf.set_line_width(0.2)
        
        for w, lines, align in zip(col_widths, all_lines, aligns):
            pdf.rect(cur_x, cur_y, w, row_h, 'DF' if fill else 'D')
            text_y = cur_y + 1.2
            for line in lines:
                pdf.set_xy(cur_x + 1, text_y)
                pdf.cell(w - 2, line_h, line, align=align)
                text_y += line_h
            cur_x += w
        pdf.set_xy(10, cur_y + row_h)
        
    if pdf.get_y() + 20 > pdf.h - 10:
        page_count += 1
        render_landscape_header(pdf, p_num=page_count)
        
    pdf.ln(3)
    v_y = pdf.get_y()
    pdf.set_fill_color(245, 247, 250)
    pdf.set_draw_color(0, 9, 83)
    pdf.set_line_width(0.4)
    pdf.rect(10, v_y, 277, 16, 'DF')
    
    pdf.set_xy(14, v_y + 1.8)
    pdf.set_font('Helvetica', 'B', 8)
    pdf.set_text_color(0, 9, 83)
    pdf.cell(270, 3.8, 'END-OF-TERM RECORD OF WORK VERIFICATION & CURRICULUM AUDIT TRAIL', ln=1)
    
    pdf.set_font('Helvetica', '', 7.5)
    pdf.set_text_color(50, 50, 50)
    pdf.set_x(14)
    pdf.cell(90, 3.5, clean_pdf_text(f'Compiled & Certified By: {trainer.upper()} (Trainer)'), ln=0)
    pdf.cell(95, 3.5, 'HOD Review & Approval: _____________________________', ln=0)
    pdf.cell(85, 3.5, 'Quality Assurance Seal & Date: ___________________', ln=1)
    
    pdf.set_x(14)
    pdf.cell(90, 3.5, clean_pdf_text(f'Signature: __________________ Date: {payload.get("compiled_date", "October 2026")}'), ln=0)
    pdf.cell(95, 3.5, 'Signature: __________________ Date: ___________________', ln=0)
    pdf.cell(85, 3.5, 'Remarks: [  ] Approved  [  ] Revisions Required', ln=1)
    
    out = pdf.output(dest='S')
    pdf_bytes = out.encode('latin1') if isinstance(out, str) else bytes(out)
    return pdf_bytes

# Helper: Parse unit level from unit code
def get_level_from_code(code: str) -> int:
    match = re.search(r'/(\d)/', code)
    if match:
        return int(match.group(1))
    parts = code.split("/")
    for p in parts:
        try:
            val = int(p)
            if 3 <= val <= 8:
                return val
        except ValueError:
            continue
    return 6  # Default fallback

# Helper: Rule-based AI exam builder matching CDACC Level rules
def build_cdacc_exam(unit_code: str, course_name: str, level: int, topics: list, outcomes: list):
    title = f"{course_name} — End of Unit Assessment"
    
    questions_a = []
    questions_b = []
    
    import re
    def is_raw_question(text):
        return bool(re.search(r'\(\d+\s*marks?\)', text, re.IGNORECASE) or '?' in text or re.match(r'^\s*(?:[a-zA-Z]\)|\d+\.)', text))
        
    def format_q(text, fallback):
        if is_raw_question(text):
            return text
        return fallback.format(text)

    # Fallback lists to ensure we have enough items
    if not outcomes:
        outcomes = ["Analyze core operational principles", "Implement standard field procedures"]
    if not topics:
        topics = ["Core Principles", "Standard Operations"]
            
    if level in [5, 6]:
        # Section A: exactly 10 questions, each 4 marks (Total = 40)
        for i in range(10):
            outcome = outcomes[i % len(outcomes)]
            q_text = format_q(outcome, "Discuss the critical parameters regarding: '{}'. Explain its structural relevance under CDACC standards.")
            questions_a.append({
                "id": f"q-sa-{i+1}",
                "text": q_text,
                "type": "short_answer",
                "correct_answer": "Grading criteria: Clarity of definition, logical architecture.",
                "marks": 4,
                "critical_aspect": f"Demonstrates comprehensive mastery of {outcome[:20].lower()}..."
            })
            
        # Section B: exactly 4 questions (Choose 3), each 20 marks (Total = 60)
        for i in range(4):
            topic = topics[i % len(topics)]
            if is_raw_question(topic):
                q_text = topic
            else:
                q_text = (
                    f"a) Formulate a complete methodology targeting '{topic}'. (8 Marks)\n"
                    f"b) Outline equipment setup and error margins associated with this process. (6 Marks)\n"
                    f"c) Explain three procedures for managing risks associated with the process. (6 Marks)"
                )
            questions_b.append({
                "id": f"q-ex-{i+1}",
                "text": q_text,
                "type": "short_answer",
                "correct_answer": "Grading rubrics: Methodology (8m), Equipment (6m), Risk Management (6m).",
                "marks": 20,
                "critical_aspect": f"Demonstrated capability for {topic[:20].lower()}..."
            })
            
    elif level == 4:
        # Section A: exactly 5 MCQs, each 2 marks (Total = 10) - Strictly no True/False
        for i in range(5):
            topic = topics[i % len(topics)]
            is_raw = is_raw_question(topic)
            q_text = format_q(topic, "Which of the following describes the standard best practice for managing '{}' under Level 4 guidelines?")
            questions_a.append({
                "id": f"q-mcq-{i+1}",
                "text": q_text,
                "type": "short_answer",
                "correct_answer": 1 if not is_raw else "",
                "marks": 2,
                "critical_aspect": f"Selects standard practices for {topic[:20].lower()}."
            })
            
        # Section B: exactly 3 questions (Choose 2), each 20 marks (Total = 40)
        for i in range(3):
            topic = topics[i % len(topics)]
            if is_raw_question(topic):
                q_text = topic
            else:
                q_text = (
                    f"a) Describe the core requirements for setting up a standardized test case for '{topic}'. (12 Marks)\n"
                    f"b) Explain the validation criteria used to measure success in this process. (8 Marks)"
                )
            questions_b.append({
                "id": f"q-sb-{i+1}",
                "text": q_text,
                "type": "short_answer",
                "correct_answer": "a) Setup requirements (12m). b) Validation criteria (8m).",
                "marks": 20,
                "critical_aspect": f"Formulates test cases for {topic[:20].lower()}."
            })
            
    else:  # Level 3
        # Section A: exactly 5 questions, each 2 marks (Total = 10)
        for i in range(5):
            topic = topics[i % len(topics)]
            is_raw = is_raw_question(topic)
            q_text = format_q(topic, "Explain the baseline element associated with '{}'.")
            questions_a.append({
                "id": f"q-mcq-{i+1}",
                "text": q_text,
                "type": "short_answer",
                "correct_answer": 0 if not is_raw else "",
                "marks": 2,
                "critical_aspect": f"Identifies baseline parameter controls for {topic[:20].lower()}."
            })
            
        # Section B: exactly 3 questions (Choose 2), each 20 marks (Total = 40)
        for i in range(3):
            topic = topics[i % len(topics)]
            if is_raw_question(topic):
                q_text = topic
            else:
                q_text = (
                    f"a) Explain the operational steps to verify and document '{topic}' values. (12 Marks)\n"
                    f"b) Identify four potential risks during this process and how to mitigate them. (8 Marks)"
                )
            questions_b.append({
                "id": f"q-sb-{i+1}",
                "text": q_text,
                "type": "short_answer",
                "correct_answer": "a) Operational steps (12m). b) Risk mitigation (8m).",
                "marks": 20,
                "critical_aspect": f"Verifies values for {topic[:20].lower()}."
            })

    payload = {
        "title": title,
        "duration_minutes": 120 if level in [5, 6] else 90,
        "total_marks": 100 if level in [5, 6] else 50,
        "instructions": "Answer all questions in Section A, and follow instructions in Section B.",
        "section_a": {
            "title": "Section A — Cognitive Theory (CT)",
            "instructions": "Answer all questions in this section.",
            "total_marks": 40 if level in [5, 6] else 10,
            "questions": questions_a
        },
        "section_b": {
            "title": "Section B — Practical Performance (CP)",
            "instructions": "Answer ANY THREE questions in this section." if level in [5, 6] else "Answer ANY TWO questions in this section.",
            "total_marks": 60 if level in [5, 6] else 40,
            "questions": questions_b
        }
    }
    return payload

def build_cdacc_practical(unit_code, course_name, level, topics, outcomes):
    title = f"FORMATIVE PRACTICAL ASSESSMENT"
    
    if not outcomes:
        outcomes = ["Analyze core operational principles", "Implement standard field procedures"]
    if not topics:
        topics = ["Core Principles", "Standard Operations"]
        
    tasks = []
    rubric = []
    
    total_marks = 100 if level in [5, 6] else 60
    
    # Generate 2 tasks
    for i in range(2):
        topic = topics[i % len(topics)]
        outcome = outcomes[i % len(outcomes)]
        
        task_marks = total_marks // 2
        tasks.append({
            "id": f"task-{i+1}",
            "title": f"TASK {i+1}: Practical implementation of {topic}",
            "marks": task_marks,
            "details": [
                f"Preparation: Ensure all tools and software are correctly configured for {topic[:30]}.",
                f"Execution: Demonstrate the ability to {outcome[:50].lower()}.",
                f"Output: Produce a verifiable output meeting industry standards."
            ]
        })
        
        rubric.append({
            "category": f"Execution of Task {i+1}",
            "max_marks": task_marks,
            "criteria": [
                { "description": f"Appropriate setup and tool selection for {topic[:20]}.", "marks": round(task_marks * 0.3) },
                { "description": f"Successful demonstration of {outcome[:30]}.", "marks": round(task_marks * 0.5) },
                { "description": "Quality and compliance of the final output.", "marks": round(task_marks * 0.2) }
            ]
        })
        
    payload = {
        "title": title,
        "type": "practical",
        "duration_minutes": 180,
        "total_marks": total_marks,
        "instructions": "Ensure all required resources for the assessment have been provided before beginning. The assessor will evaluate you dynamically as you perform the tasks.",
        "project_brief": f"In this project, the candidate will demonstrate competence in {course_name} by completing hands-on scenarios based on standard industry requirements.",
        "elements_covered": topics[:4],
        "tasks": tasks,
        "rubric": rubric
    }
    return payload


# Canonical standard mark breakdowns for Digital Literacy (061155101A-WA1)
DEFAULT_BREAKDOWNS = {
    1: [
        {"criterion": "Definition of digital literacy (access, evaluate, create info)", "marks": 1},
        {"criterion": "Importance in modern workplace (efficiency, automation, communication)", "marks": 1}
    ],
    2: [
        {"criterion": "Input device 1 (Keyboard) and Input device 2 (Optical Mouse)", "marks": 2},
        {"criterion": "Output device 1 (Display Monitor) and Output device 2 (Laser Printer)", "marks": 2}
    ],
    3: [
        {"criterion": "Step 1: Saving active work and open files", "marks": 1},
        {"criterion": "Step 2: Gracefully closing active software applications", "marks": 1},
        {"criterion": "Step 3: Initiating Start > Power > Shut Down command sequence", "marks": 1}
    ],
    4: [
        {"criterion": "Technique 1: Touch typing utilizing home row keys (ASDF JKL;)", "marks": 1},
        {"criterion": "Technique 2: Effective application of keyboard shortcuts (Ctrl+C, Ctrl+V)", "marks": 1}
    ],
    5: [
        {"criterion": "Step 1: New folder creation procedure (Right click > New > Folder)", "marks": 1},
        {"criterion": "Step 2: Saving confidential report into target folder (File > Save As)", "marks": 1},
        {"criterion": "Step 3: Setting password protection / document encryption", "marks": 2}
    ],
    6: [
        {"criterion": "Area 1: Healthcare / Hospitals (Electronic medical records & diagnostics)", "marks": 1},
        {"criterion": "Area 2: Banking & Finance (ATMs, ledger accounting & fund transfers)", "marks": 1},
        {"criterion": "Area 3: Educational Institutions (Computer-assisted learning & administration)", "marks": 1}
    ],
    7: [
        {"criterion": "Step 1: Physical video cable connection (HDMI/VGA) and power on", "marks": 1},
        {"criterion": "Step 2: Display projection mode selection (Windows Key + P > Duplicate)", "marks": 1},
        {"criterion": "Step 3: Screen calibration and projector lens focus adjustment", "marks": 1}
    ],
    8: [
        {"criterion": "Conceptual distinction (System software manages hardware vs Application serves user tasks)", "marks": 2},
        {"criterion": "Two system software examples (e.g. Windows 11, Linux)", "marks": 1},
        {"criterion": "Two application software examples (e.g. MS Word, MS Excel)", "marks": 1}
    ],
    9: [
        {"criterion": "Action definition (Click, hold, move cursor across screen, release)", "marks": 1},
        {"criterion": "File management application (Relocating or copying files/folders)", "marks": 1}
    ],
    10: [
        {"criterion": "Policy 1: Acceptable Use Policy (AUP) regulating internet and device activity", "marks": 1.5},
        {"criterion": "Policy 2: Password Security & Data Protection rules safeguarding confidentiality", "marks": 1.5}
    ],
    11: [
        {"criterion": "11(a) Word processor primary function & 2 differentiating features from plain text", "marks": 4},
        {"criterion": "11(b) Four newsletter formatting options (columns, drop caps, borders, typography)", "marks": 4},
        {"criterion": "11(c) Key combination functions: i. CTRL+Z, ii. CTRL+X, iii. CTRL+V", "marks": 6},
        {"criterion": "11(d) Three technical differences between a calculator and a computer", "marks": 6}
    ],
    12: [
        {"criterion": "12(a) Procedure to insert and format a pre-designed Word table for sales data", "marks": 6},
        {"criterion": "12(b) Steps to save spreadsheet to specific network drive and print page 1 only", "marks": 4},
        {"criterion": "12(c) Four essential components of a spreadsheet used to build a budget", "marks": 4},
        {"criterion": "12(d) Three lab safety measures to protect computer equipment and personnel", "marks": 6}
    ],
    13: [
        {"criterion": "13(a) Difference between a formula and a function in Excel with examples", "marks": 6},
        {"criterion": "13(b) Steps to sort customer records alphabetically from A to Z", "marks": 4},
        {"criterion": "13(c) Word processing features to design a professional two-page CV", "marks": 6},
        {"criterion": "13(d) Four technical factors to consider when procuring an Operating System", "marks": 4}
    ]
}

DEFAULT_EXAM_QUESTIONS = {
    1: {
        "text": "Define digital literacy and state its importance in a modern workplace.",
        "max_marks": 2,
        "breakdown": DEFAULT_BREAKDOWNS[1]
    },
    2: {
        "text": "List two input devices and two output devices commonly used in a professional setting.",
        "max_marks": 4,
        "breakdown": DEFAULT_BREAKDOWNS[2]
    },
    3: {
        "text": "Explain the correct sequence of steps to safely shut down a computer device as per workplace procedure.",
        "max_marks": 3,
        "breakdown": DEFAULT_BREAKDOWNS[3]
    },
    4: {
        "text": "Describe two keyboard techniques that enhance efficiency when typing a report.",
        "max_marks": 2,
        "breakdown": DEFAULT_BREAKDOWNS[4]
    },
    5: {
        "text": "You have just completed drafting a confidential report. Outline the steps you would take to create a new folder, save the report into it, and then password protect the folder, ensuring data security.",
        "max_marks": 4,
        "breakdown": DEFAULT_BREAKDOWNS[5]
    },
    6: {
        "text": "Identify three common areas where computers are used",
        "max_marks": 3,
        "breakdown": DEFAULT_BREAKDOWNS[6]
    },
    7: {
        "text": "A colleague needs to present a document from their laptop to a large screen. Describe the steps to connect an external projector to a laptop.",
        "max_marks": 3,
        "breakdown": DEFAULT_BREAKDOWNS[7]
    },
    8: {
        "text": "Differentiate between System Software and Application Software, providing two examples for each.",
        "max_marks": 4,
        "breakdown": DEFAULT_BREAKDOWNS[8]
    },
    9: {
        "text": "Explain the concept of 'drag and drop' in the context of file management",
        "max_marks": 2,
        "breakdown": DEFAULT_BREAKDOWNS[9]
    },
    10: {
        "text": "State two workplace policies or regulations that guide the operation of computer devices.",
        "max_marks": 3,
        "breakdown": DEFAULT_BREAKDOWNS[10]
    },
    11: {
        "text": "11.\na) Describe the primary function of a word processing application and list two features that differentiate it from a plain text editor. (4 Marks)\nb) You are tasked with creating a newsletter in a word processing application. List four formatting options you would apply to make it visually appealing and easy to read. (4 Marks)\nc) What is the function of the following key combinations in Ms Office? (6 Marks)\n   i. CTRL + Z\n   ii. CTRL + X\n   iii. CTRL + V\nd) Describe three differences between a calculator and a computer (6 Marks)",
        "max_marks": 20,
        "breakdown": DEFAULT_BREAKDOWNS[11]
    },
    12: {
        "text": "12.\na) Explain the steps to insert a pre-designed table into a Word document to organize monthly sales data. (6 Marks)\nb) You have completed a critical financial report in a spreadsheet. Describe how you would save it to a specific network drive and then print only the first page. (4 Marks)\nc) A manager asks you to create a simple budget in a spreadsheet application. List four essential components of a spreadsheet you would use to build and organize this data. (4 Marks)\nd) When using the computer laboratory, several measures can be put in place to ensure computers are safe. Explain three measures. (6 Marks)",
        "max_marks": 20,
        "breakdown": DEFAULT_BREAKDOWNS[12]
    },
    13: {
        "text": "13.\na) State the difference between a formula and a function in a spreadsheet application, giving an example of each. (6 Marks)\nb) Explain the procedure to sort customer records alphabetically from A to Z in a spreadsheet. (4 Marks)\nc) Highlight four word processing features that would be essential in designing a professional two-page curriculum vitae (CV). (6 Marks)\nd) Identify four technical factors to consider when choosing an Operating System for an office environment. (4 Marks)",
        "max_marks": 20,
        "breakdown": DEFAULT_BREAKDOWNS[13]
    }
}


def generate_exam_results_docx(payload):
    doc = docx.Document()
    section = doc.sections[0]
    section.top_margin = docx.shared.Inches(0.5)
    section.bottom_margin = docx.shared.Inches(0.5)
    section.left_margin = docx.shared.Inches(0.7)
    section.right_margin = docx.shared.Inches(0.7)

    style = doc.styles['Normal']
    style.font.name = 'Maiandra GD'
    style.font.size = docx.shared.Pt(9.5)

    current_dir = os.path.dirname(os.path.abspath(__file__))
    project_root = os.path.dirname(current_dir)
    logo_path = os.path.join(project_root, "client", "public", "mtti-logo.jpg")
    if not os.path.exists(logo_path):
        logo_path = os.path.abspath("client/public/mtti-logo.jpg")

    if os.path.exists(logo_path):
        p_logo = doc.add_paragraph()
        p_logo.alignment = 1
        p_logo.paragraph_format.space_after = docx.shared.Pt(2)
        try:
            p_logo.add_run().add_picture(logo_path, width=docx.shared.Inches(0.95))
        except Exception:
            pass

    p_hdr = doc.add_paragraph()
    p_hdr.alignment = 1
    p_hdr.paragraph_format.space_after = docx.shared.Pt(3)

    r1 = p_hdr.add_run("MUKIRIA TECHNICAL TRAINING INSTITUTE\n")
    r1.bold = True
    r1.font.size = docx.shared.Pt(13)
    try:
        r1.font.color.rgb = docx.shared.RGBColor(0, 9, 83)
    except Exception:
        pass

    dept_str = str(payload.get('department') or "HOSPITALITY DEPARTMENT\nBUILDING DEPARTMENT").upper()
    for d_line in dept_str.split("\n"):
        if d_line.strip():
            r_d = p_hdr.add_run(f"{d_line.strip()}\n")
            r_d.bold = True
            r_d.font.size = docx.shared.Pt(10)

    r2 = p_hdr.add_run("INTERNAL EXAMINATION\n")
    r2.bold = True
    r2.font.size = docx.shared.Pt(10)

    exam_title = str(payload.get('exam_title', 'WRITTEN ASSESSMENT 1')).upper()
    r3 = p_hdr.add_run(f"{exam_title}\n")
    r3.bold = True
    r3.font.size = docx.shared.Pt(11)
    try:
        r3.font.color.rgb = docx.shared.RGBColor(196, 136, 32)
    except Exception:
        pass

    time_allowed = str(payload.get('time_allowed', 'Time: 2 HOURS'))
    if not time_allowed.lower().startswith('time'):
        time_allowed = f"Time: {time_allowed}"
    r4 = p_hdr.add_run(f"{time_allowed}\n")
    r4.bold = True
    r4.font.size = docx.shared.Pt(9.5)
    try:
        r4.font.color.rgb = docx.shared.RGBColor(0, 9, 83)
    except Exception:
        pass

    # Metadata Paragraph
    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.line_spacing = 1.15
    p_meta.paragraph_format.space_after = docx.shared.Pt(4)

    metadata = [
        ("COURSE NAME", payload.get("course_name") or "OFFICE ADMINISTRATION LEVEL 5 & 6, LAND SURVEY LEVEL 5 & 6"),
        ("COURSE CODE", payload.get("course_code") or payload.get("unit_code") or "061155101A"),
        ("UNIT NAME", payload.get("unit_name") or "APPLY DIGITAL LITERACY"),
        ("CLASS", payload.get("class_code") or "FBS5/6/J/25, LS5/6/S/25"),
        ("SERIES", payload.get("series") or "SEP - NOV 2026"),
    ]
    for label, val in metadata:
        r_lbl = p_meta.add_run(f"{label}: ")
        r_lbl.bold = True
        r_lbl.font.size = docx.shared.Pt(8.5)
        r_val = p_meta.add_run(f"{val}\n")
        r_val.font.size = docx.shared.Pt(8.5)

    # Red Pen Score Tally Stamp
    raw_total = payload.get("total_score", 0)
    try:
        total_score = round(float(raw_total), 1)
        if total_score.is_integer():
            total_score = int(total_score)
    except Exception:
        total_score = raw_total
    total_marks = payload.get("total_marks", 70)
    pct = payload.get("percentage")
    if pct is None:
        pct = round((float(total_score) / float(total_marks) * 100) if total_marks else 0)
    grade = str(payload.get("grade", "Pass"))
    status = str(payload.get("status", "COMPETENT (PASS)" if pct >= 50 else "NOT YET COMPETENT (REFER)"))

    sec_a_items = payload.get("section_a") or payload.get("answers", {}).get("section_a", [])
    sec_b_items = payload.get("section_b") or payload.get("answers", {}).get("section_b", [])
    sec_a_subtotal = sum(float(it.get("marks_awarded", 0)) for it in sec_a_items)
    if sec_a_subtotal.is_integer():
        sec_a_subtotal = int(sec_a_subtotal)
    sec_b_subtotal = sum(float(it.get("marks_awarded", 0)) for it in sec_b_items)
    if sec_b_subtotal.is_integer():
        sec_b_subtotal = int(sec_b_subtotal)

    cand_name = str(payload.get('student_name', 'Alex Kinoti'))
    cand_reg = str(payload.get('reg_number', '10525'))
    trainer_n = str(payload.get('trainer_name', 'Alexander Kinoti'))
    eval_date = str(payload.get('evaluated_at', '07/10/2026'))

    tbl_stamp = doc.add_table(rows=3, cols=2)
    tbl_stamp.style = 'Table Grid'
    # Row 0: Banner
    cell_top = tbl_stamp.rows[0].cells[0]
    tbl_stamp.rows[0].cells[1].merge(cell_top)
    p_bann = cell_top.paragraphs[0]
    p_bann.alignment = 1
    r_bann = p_bann.add_run("OFFICIAL EVALUATION SCORE TALLY -- SIMULATED RED PEN ASSESSMENT")
    r_bann.bold = True
    r_bann.font.size = docx.shared.Pt(8)
    try:
        r_bann.font.color.rgb = docx.shared.RGBColor(185, 28, 28)
    except Exception:
        pass

    # Row 1: Candidate & Assessor Details
    c_r1_1 = tbl_stamp.rows[1].cells[0].paragraphs[0]
    c_r1_1.add_run(f"Candidate: {cand_name} (Reg: {cand_reg})").font.size = docx.shared.Pt(8)
    c_r1_2 = tbl_stamp.rows[1].cells[1].paragraphs[0]
    c_r1_2.add_run(f"Assessor: {trainer_n} | Date: {eval_date}").font.size = docx.shared.Pt(8)

    # Row 2: Score Summary
    cell_bot = tbl_stamp.rows[2].cells[0]
    tbl_stamp.rows[2].cells[1].merge(cell_bot)
    p_sc = cell_bot.paragraphs[0]
    p_sc.alignment = 1
    r_sc = p_sc.add_run(f"SECTION A: {sec_a_subtotal}/30  |  SECTION B: {sec_b_subtotal}/40  |  TOTAL: {total_score}/{total_marks} ({pct}%)  --  {status}")
    r_sc.bold = True
    r_sc.font.size = docx.shared.Pt(10)
    try:
        r_sc.font.color.rgb = docx.shared.RGBColor(220, 38, 38)
    except Exception:
        pass

    for row in tbl_stamp.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_before = docx.shared.Pt(2)
                p.paragraph_format.space_after = docx.shared.Pt(2)

    # Instructions to Candidate
    p_inst = doc.add_paragraph()
    p_inst.paragraph_format.space_before = docx.shared.Pt(6)
    p_inst.paragraph_format.space_after = docx.shared.Pt(2)
    r_ih = p_inst.add_run("INSTRUCTIONS TO CANDIDATE")
    r_ih.bold = True
    r_ih.font.size = docx.shared.Pt(9)

    instructions = [
        "1. This paper consists of two sections A and B;",
        "2. Answer ALL the question as guided in each section;",
        "3. Marks for each question are as indicated in the brackets;",
        "4. You are provided with a separate answer booklet to answer the questions;",
        "5. Do not write in this question paper."
    ]
    for inst in instructions:
        p_i = doc.add_paragraph(inst)
        p_i.paragraph_format.space_after = docx.shared.Pt(1)
        if p_i.runs:
            p_i.runs[0].font.size = docx.shared.Pt(8)

    def render_docx_question(item, default_num, is_sec_b=False):
        q_num = item.get("q_num", default_num)
        default_meta = DEFAULT_EXAM_QUESTIONS.get(q_num, {})
        q_text_raw = item.get("text")
        if not q_text_raw or q_text_raw.lower() in [f"question {q_num}", f"task {q_num}", f"question {default_num}", f"task {default_num}"]:
            q_text_raw = default_meta.get("text", f"Question {q_num}")
        q_text = str(q_text_raw)

        raw_awarded = item.get("marks_awarded", 0)
        try:
            awarded = round(float(raw_awarded), 1)
            if awarded.is_integer():
                awarded = int(awarded)
        except Exception:
            awarded = raw_awarded
        max_m = item.get("max_marks", default_meta.get("max_marks", 20 if is_sec_b else 2))
        ans_text = str(item.get("answer", ""))
        reasoning = str(item.get("ai_reasoning") or item.get("notes") or "")

        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = docx.shared.Pt(5)
        p_q.paragraph_format.space_after = docx.shared.Pt(1)
        q_hdr = f"{q_num}. {q_text}" if not q_text.startswith(f"{q_num}.") else q_text
        if not f"({max_m}" in q_hdr and not f"{max_m} Mark" in q_hdr:
            q_hdr += f" ({max_m} Marks)"
        r_q = p_q.add_run(q_hdr)
        r_q.bold = True
        r_q.font.size = docx.shared.Pt(8.5)

        # Candidate response
        p_ans = doc.add_paragraph()
        p_ans.paragraph_format.left_indent = docx.shared.Inches(0.25)
        p_ans.paragraph_format.space_after = docx.shared.Pt(1)
        r_ans_lbl = p_ans.add_run("Candidate's Written Response:\n")
        r_ans_lbl.bold = True
        r_ans_lbl.font.size = docx.shared.Pt(7.5)
        try:
            r_ans_lbl.font.color.rgb = docx.shared.RGBColor(80, 80, 80)
        except Exception:
            pass
        r_ans_val = p_ans.add_run(ans_text if ans_text.strip() else "[No candidate response entered]")
        r_ans_val.font.size = docx.shared.Pt(8)

        # Rubric breakdown
        breakdown = item.get("breakdown") or default_meta.get("breakdown") or DEFAULT_BREAKDOWNS.get(q_num, [])
        if breakdown:
            p_rb_hdr = doc.add_paragraph()
            p_rb_hdr.paragraph_format.left_indent = docx.shared.Inches(0.25)
            p_rb_hdr.paragraph_format.space_after = docx.shared.Pt(1)
            r_rb_lbl = p_rb_hdr.add_run("Official Assessor Mark Distribution & Rubric Breakdown:")
            r_rb_lbl.bold = True
            r_rb_lbl.font.size = docx.shared.Pt(7.5)
            try:
                r_rb_lbl.font.color.rgb = docx.shared.RGBColor(185, 28, 28)
            except Exception:
                pass

            rem_marks = float(awarded) if isinstance(awarded, (int, float)) else 0.0
            for crit_idx, crit in enumerate(breakdown):
                c_desc = str(crit.get("criterion", f"Mark Component {crit_idx+1}"))
                c_max = float(crit.get("marks", 1))
                if isinstance(awarded, (int, float)):
                    if awarded == max_m:
                        c_awarded = c_max
                    elif awarded == 0:
                        c_awarded = 0.0
                    else:
                        c_awarded = min(c_max, max(0.0, rem_marks))
                        rem_marks -= c_awarded
                else:
                    c_awarded = 0.0

                c_awarded_disp = int(c_awarded) if c_awarded.is_integer() else round(c_awarded, 1)
                c_max_disp = int(c_max) if c_max.is_integer() else round(c_max, 1)
                is_crit_full = (c_awarded >= c_max)

                p_crit = doc.add_paragraph()
                p_crit.paragraph_format.left_indent = docx.shared.Inches(0.35)
                p_crit.paragraph_format.space_after = docx.shared.Pt(1)
                r_c_desc = p_crit.add_run(f"* {c_desc}: ")
                r_c_desc.font.size = docx.shared.Pt(7.5)
                if is_crit_full:
                    r_c_tag = p_crit.add_run(f"[Awarded {c_awarded_disp} / {c_max_disp} Mark{'s' if c_max_disp > 1 else ''}]")
                    r_c_tag.bold = True
                    r_c_tag.font.size = docx.shared.Pt(7.5)
                    try:
                        r_c_tag.font.color.rgb = docx.shared.RGBColor(34, 139, 34)
                    except Exception:
                        pass
                elif c_awarded > 0:
                    r_c_tag = p_crit.add_run(f"[Partial {c_awarded_disp} / {c_max_disp} Marks]")
                    r_c_tag.bold = True
                    r_c_tag.font.size = docx.shared.Pt(7.5)
                    try:
                        r_c_tag.font.color.rgb = docx.shared.RGBColor(217, 119, 6)
                    except Exception:
                        pass
                else:
                    r_c_tag = p_crit.add_run(f"[0 / {c_max_disp} Marks]")
                    r_c_tag.bold = True
                    r_c_tag.font.size = docx.shared.Pt(7.5)
                    try:
                        r_c_tag.font.color.rgb = docx.shared.RGBColor(220, 38, 38)
                    except Exception:
                        pass

        # Question Total Awarded Badge
        p_badge = doc.add_paragraph()
        p_badge.paragraph_format.left_indent = docx.shared.Inches(0.25)
        p_badge.paragraph_format.space_after = docx.shared.Pt(2)
        is_full = (isinstance(awarded, (int, float)) and awarded >= max_m)
        is_zero = (awarded == 0)
        verdict_tag = "FULL CREDIT" if is_full else ("NIL CREDIT" if is_zero else "PARTIAL CREDIT")
        r_bdg = p_badge.add_run(f">> Question {q_num} Total Awarded: {awarded} / {max_m} Marks [{verdict_tag}]")
        r_bdg.bold = True
        r_bdg.font.size = docx.shared.Pt(8)
        try:
            r_bdg.font.color.rgb = docx.shared.RGBColor(220, 38, 38)
        except Exception:
            pass

        if reasoning:
            p_note = doc.add_paragraph()
            p_note.paragraph_format.left_indent = docx.shared.Inches(0.35)
            p_note.paragraph_format.space_after = docx.shared.Pt(2)
            r_note = p_note.add_run(f"Assessor Note: {reasoning}")
            r_note.italic = True
            r_note.font.size = docx.shared.Pt(7.5)
            try:
                r_note.font.color.rgb = docx.shared.RGBColor(185, 28, 28)
            except Exception:
                pass

    # Section A
    p_a_hdr = doc.add_paragraph()
    p_a_hdr.paragraph_format.space_before = docx.shared.Pt(8)
    p_a_hdr.paragraph_format.space_after = docx.shared.Pt(2)
    r_ah = p_a_hdr.add_run("SECTION A (30 MARKS) -- Answer ALL Questions")
    r_ah.bold = True
    r_ah.font.size = docx.shared.Pt(10)
    try:
        r_ah.font.color.rgb = docx.shared.RGBColor(0, 9, 83)
    except Exception:
        pass
    r_asub = p_a_hdr.add_run(f"    (Subtotal: {sec_a_subtotal} / 30 Marks)")
    r_asub.bold = True
    r_asub.font.size = docx.shared.Pt(8.5)
    try:
        r_asub.font.color.rgb = docx.shared.RGBColor(220, 38, 38)
    except Exception:
        pass

    for idx, item in enumerate(sec_a_items):
        render_docx_question(item, idx + 1, is_sec_b=False)

    # Section B
    p_b_hdr = doc.add_paragraph()
    p_b_hdr.paragraph_format.space_before = docx.shared.Pt(8)
    p_b_hdr.paragraph_format.space_after = docx.shared.Pt(2)
    r_bh = p_b_hdr.add_run("SECTION B (40 MARKS) -- Answer ANY TWO Questions")
    r_bh.bold = True
    r_bh.font.size = docx.shared.Pt(10)
    try:
        r_bh.font.color.rgb = docx.shared.RGBColor(0, 9, 83)
    except Exception:
        pass
    r_bsub = p_b_hdr.add_run(f"    (Subtotal: {sec_b_subtotal} / 40 Marks)")
    r_bsub.bold = True
    r_bsub.font.size = docx.shared.Pt(8.5)
    try:
        r_bsub.font.color.rgb = docx.shared.RGBColor(220, 38, 38)
    except Exception:
        pass

    for idx, item in enumerate(sec_b_items):
        render_docx_question(item, idx + 11, is_sec_b=True)

    # Overall Remarks
    p_rem_hdr = doc.add_paragraph()
    p_rem_hdr.paragraph_format.space_before = docx.shared.Pt(8)
    p_rem_hdr.paragraph_format.space_after = docx.shared.Pt(2)
    r_rh = p_rem_hdr.add_run("ASSESSOR / TRAINER EXAMINATION FEEDBACK REMARKS:")
    r_rh.bold = True
    r_rh.font.size = docx.shared.Pt(8.5)
    try:
        r_rh.font.color.rgb = docx.shared.RGBColor(0, 9, 83)
    except Exception:
        pass

    comments = str(payload.get("trainer_comments") or "The candidate demonstrates exceptional competence in digital literacy principles, practical workplace computer procedures, and software applications.")
    p_rem_txt = doc.add_paragraph(f'"{comments}"')
    p_rem_txt.paragraph_format.left_indent = docx.shared.Inches(0.25)
    p_rem_txt.paragraph_format.space_after = docx.shared.Pt(6)
    if p_rem_txt.runs:
        p_rem_txt.runs[0].italic = True
        p_rem_txt.runs[0].font.size = docx.shared.Pt(8)

    # Signatures Table
    tbl_sig = doc.add_table(rows=3, cols=2)
    tbl_sig.style = 'Table Grid'
    tbl_sig.rows[0].cells[0].text = f"Assessor / Trainer: {trainer_n}"
    tbl_sig.rows[0].cells[1].text = "Internal Verifier / HOD: Computing & Informatics"
    tbl_sig.rows[1].cells[0].text = f"Signature: {trainer_n} (Digitally Certified Assessor)"
    tbl_sig.rows[1].cells[1].text = "Signature: ____________________________________"
    tbl_sig.rows[2].cells[0].text = f"Date: {eval_date}"
    tbl_sig.rows[2].cells[1].text = "Official Stamp: Mukiria Assessment Center 01200004"
    for row in tbl_sig.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_before = docx.shared.Pt(2)
                p.paragraph_format.space_after = docx.shared.Pt(2)
                if p.runs:
                    p.runs[0].font.size = docx.shared.Pt(8)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


class MukiriaMarkedExamPDF(FPDF):
    def __init__(self):
        super().__init__(orientation='P', unit='mm', format='A4')
        self.alias_nb_pages()
        self.set_margins(14, 12, 14)
        self.set_auto_page_break(True, margin=15)

    def header(self):
        self.set_font('Helvetica', '', 8)
        self.set_text_color(100, 100, 100)
        self.cell(0, 5, f'Page {self.page_no()} of {{nb}}', 0, 1, 'L')
        self.ln(1)

    def footer(self):
        self.set_y(-12)
        self.set_font('Helvetica', 'I', 7)
        self.set_text_color(120, 120, 120)
        self.cell(0, 5, 'Mukiria Technical Training Institute -- Official Evaluated Examination Script', 0, 0, 'C')


def generate_exam_results_pdf(payload):
    pdf = MukiriaMarkedExamPDF()
    pdf.add_page()

    current_dir = os.path.dirname(os.path.abspath(__file__))
    project_root = os.path.dirname(current_dir)
    logo_path = os.path.join(project_root, "client", "public", "mtti-logo.jpg")
    if not os.path.exists(logo_path):
        logo_path = os.path.abspath("client/public/mtti-logo.jpg")

    if os.path.exists(logo_path):
        try:
            pdf.image(logo_path, x=95, y=pdf.get_y(), w=20)
            pdf.set_y(pdf.get_y() + 21)
        except Exception:
            pass

    # Institution branding headers matching media_1790901424074.pdf
    pdf.set_font('Helvetica', 'B', 13)
    pdf.set_text_color(0, 9, 83)
    pdf.cell(0, 5.5, 'MUKIRIA TECHNICAL TRAINING INSTITUTE', ln=1, align='C')

    pdf.set_font('Helvetica', 'B', 10)
    pdf.set_text_color(40, 40, 40)
    dept_str = payload.get('department')
    if not dept_str:
        dept_str = "HOSPITALITY DEPARTMENT\nBUILDING DEPARTMENT"
    for dept_line in dept_str.split("\n"):
        if dept_line.strip():
            pdf.cell(0, 4.8, clean_pdf_text(dept_line.strip().upper()), ln=1, align='C')

    pdf.cell(0, 4.8, 'INTERNAL EXAMINATION', ln=1, align='C')

    exam_title = clean_pdf_text(payload.get('exam_title', 'WRITTEN ASSESSMENT 1').upper())
    pdf.set_font('Helvetica', 'B', 11)
    pdf.set_text_color(196, 136, 32)
    pdf.cell(0, 5.2, exam_title, ln=1, align='C')

    time_allowed = clean_pdf_text(payload.get('time_allowed', 'Time: 2 HOURS'))
    if not time_allowed.lower().startswith('time'):
        time_allowed = f'Time: {time_allowed}'
    pdf.set_font('Helvetica', 'B', 9.5)
    pdf.set_text_color(0, 9, 83)
    pdf.cell(0, 4.8, time_allowed, ln=1, align='C')
    pdf.ln(1.5)

    # Exam metadata block (Unaltered, exactly matching uploaded question paper)
    metadata = [
        ('COURSE NAME', clean_pdf_text(payload.get('course_name') or 'OFFICE ADMINISTRATION LEVEL 5 & 6, LAND SURVEY LEVEL 5 & 6')),
        ('COURSE CODE', clean_pdf_text(payload.get('course_code') or payload.get('unit_code') or '061155101A')),
        ('UNIT NAME', clean_pdf_text(payload.get('unit_name') or 'APPLY DIGITAL LITERACY')),
        ('CLASS', clean_pdf_text(payload.get('class_code') or 'FBS5/6/J/25, LS5/6/S/25')),
        ('SERIES', clean_pdf_text(payload.get('series') or 'SEP - NOV 2026')),
    ]
    for label, val in metadata:
        pdf.set_font('Helvetica', 'B', 8.5)
        pdf.set_text_color(0, 0, 0)
        pdf.write(4.3, f'{label}: ')
        pdf.set_font('Helvetica', '', 8.5)
        pdf.write(4.3, f'{val}\n')
    pdf.ln(2)

    # Official Evaluation Score Tally Stamp Box
    raw_total = payload.get("total_score", 0)
    try:
        total_score = round(float(raw_total), 1)
        if total_score.is_integer():
            total_score = int(total_score)
    except Exception:
        total_score = raw_total
    total_marks = payload.get("total_marks", 70)
    pct = payload.get("percentage")
    if pct is None:
        pct = round((float(total_score) / float(total_marks) * 100) if total_marks else 0)
    grade = clean_pdf_text(str(payload.get("grade", "Pass")))
    status = clean_pdf_text(str(payload.get("status", "COMPETENT (PASS)" if pct >= 50 else "NOT YET COMPETENT (REFER)")))

    box_y = pdf.get_y()
    pdf.set_fill_color(254, 242, 242)
    pdf.set_draw_color(220, 38, 38)
    pdf.set_line_width(0.6)
    pdf.rect(14, box_y, 182, 22, 'DF')

    pdf.set_xy(16, box_y + 1.2)
    pdf.set_font('Helvetica', 'B', 8)
    pdf.set_text_color(185, 28, 28)
    pdf.cell(178, 3.5, "OFFICIAL EVALUATION SCORE TALLY -- SIMULATED RED PEN ASSESSMENT", ln=1)

    pdf.set_xy(16, box_y + 5.5)
    pdf.set_font('Helvetica', 'B', 8)
    pdf.set_text_color(20, 20, 20)
    cand_name = clean_pdf_text(payload.get('student_name', 'Alex Kinoti'))
    cand_reg = clean_pdf_text(payload.get('reg_number', '10525'))
    pdf.cell(90, 4, f"Candidate Name: {cand_name}", ln=0)
    pdf.cell(88, 4, f"Admission / Reg No: {cand_reg}", ln=1)

    sec_a_items = payload.get("section_a") or payload.get("answers", {}).get("section_a", [])
    sec_b_items = payload.get("section_b") or payload.get("answers", {}).get("section_b", [])
    sec_a_subtotal = sum(float(it.get("marks_awarded", 0)) for it in sec_a_items)
    if sec_a_subtotal.is_integer():
        sec_a_subtotal = int(sec_a_subtotal)
    sec_b_subtotal = sum(float(it.get("marks_awarded", 0)) for it in sec_b_items)
    if sec_b_subtotal.is_integer():
        sec_b_subtotal = int(sec_b_subtotal)

    pdf.set_xy(16, box_y + 10)
    pdf.set_font('Helvetica', 'B', 8)
    pdf.set_text_color(20, 20, 20)
    trainer_n = clean_pdf_text(payload.get('trainer_name', 'Alexander Kinoti'))
    eval_date = clean_pdf_text(payload.get('evaluated_at', '07/10/2026'))
    pdf.cell(90, 4, f"Assessor / Trainer: {trainer_n}", ln=0)
    pdf.cell(88, 4, f"Date Evaluated: {eval_date}", ln=1)

    pdf.set_xy(16, box_y + 15)
    pdf.set_font('Helvetica', 'B', 9.5)
    pdf.set_text_color(220, 38, 38)
    score_summary = f"SECTION A: {sec_a_subtotal}/30  |  SECTION B: {sec_b_subtotal}/40  |  TOTAL: {total_score}/{total_marks} ({pct}%)  --  {status}"
    pdf.cell(178, 5, score_summary, ln=1)

    pdf.set_y(box_y + 24)

    # Instructions to candidate
    pdf.set_font('Helvetica', 'B', 9)
    pdf.set_text_color(0, 0, 0)
    pdf.cell(0, 4.5, "INSTRUCTIONS TO CANDIDATE", ln=1)

    instructions = [
        "1. This paper consists of two sections A and B;",
        "2. Answer ALL the question as guided in each section;",
        "3. Marks for each question are as indicated in the brackets;",
        "4. You are provided with a separate answer booklet to answer the questions;",
        "5. Do not write in this question paper."
    ]
    pdf.set_font('Helvetica', '', 8)
    pdf.set_text_color(30, 30, 30)
    for inst in instructions:
        pdf.cell(0, 3.8, clean_pdf_text(inst), ln=1)

    pdf.set_font('Helvetica', 'B', 8)
    pdf.cell(0, 4.2, "THIS PAPER CONSISTS OF {nb} PRINTED PAGES", ln=1)
    pdf.ln(2)

    # Helper function to render a question with student answer and distributed marks
    def render_marked_question(item, default_num, is_section_b=False):
        q_num = item.get("q_num", default_num)
        default_meta = DEFAULT_EXAM_QUESTIONS.get(q_num, {})
        q_text_raw = item.get("text")
        if not q_text_raw or q_text_raw.lower() in [f"question {q_num}", f"task {q_num}", f"question {default_num}", f"task {default_num}"]:
            q_text_raw = default_meta.get("text", f"Question {q_num}")
        q_text = clean_pdf_text(str(q_text_raw))

        raw_awarded = item.get("marks_awarded", 0)
        try:
            awarded = round(float(raw_awarded), 1)
            if awarded.is_integer():
                awarded = int(awarded)
        except Exception:
            awarded = raw_awarded
        max_m = item.get("max_marks", default_meta.get("max_marks", 20 if is_section_b else 2))
        ans_text = clean_pdf_text(str(item.get("answer", "")))
        reasoning = clean_pdf_text(str(item.get("ai_reasoning") or item.get("notes") or ""))

        # Check page break margin
        if pdf.get_y() > 235:
            pdf.add_page()

        # Question Header (Unaltered from exam paper)
        pdf.set_font('Helvetica', 'B', 8.5)
        pdf.set_text_color(0, 0, 0)
        q_header = f"{q_num}. {q_text}" if not q_text.startswith(f"{q_num}.") else q_text
        if not f"({max_m}" in q_header and not f"{max_m} Mark" in q_header:
            q_header += f" ({max_m} Marks)"
        pdf.multi_cell(182, 4.2, q_header)

        # Candidate's Written Response block
        pdf.set_font('Helvetica', 'B', 7.5)
        pdf.set_text_color(80, 80, 80)
        pdf.set_x(18)
        pdf.cell(178, 3.5, "Candidate's Written Response:", ln=1)

        pdf.set_font('Helvetica', '', 8)
        pdf.set_text_color(30, 30, 30)
        ans_display = ans_text if ans_text.strip() else "[No candidate response entered]"
        pdf.set_x(18)
        pdf.multi_cell(176, 3.8, ans_display)

        # Assessor Mark Distribution & Rubric Breakdown block
        breakdown = item.get("breakdown") or default_meta.get("breakdown") or DEFAULT_BREAKDOWNS.get(q_num, [])
        pdf.set_x(18)
        pdf.set_font('Helvetica', 'B', 7.5)
        pdf.set_text_color(185, 28, 28)
        pdf.cell(178, 3.5, "Official Assessor Mark Distribution & Rubric Breakdown:", ln=1)

        # Distribute marks across criteria
        if breakdown:
            rem_marks = float(awarded) if isinstance(awarded, (int, float)) else 0.0
            for crit_idx, crit in enumerate(breakdown):
                c_desc = clean_pdf_text(crit.get("criterion", f"Mark Component {crit_idx+1}"))
                c_max = float(crit.get("marks", 1))
                if isinstance(awarded, (int, float)):
                    if awarded == max_m:
                        c_awarded = c_max
                    elif awarded == 0:
                        c_awarded = 0.0
                    else:
                        c_awarded = min(c_max, max(0.0, rem_marks))
                        rem_marks -= c_awarded
                else:
                    c_awarded = 0.0

                c_awarded_disp = int(c_awarded) if c_awarded.is_integer() else round(c_awarded, 1)
                c_max_disp = int(c_max) if c_max.is_integer() else round(c_max, 1)
                is_crit_full = (c_awarded >= c_max)

                pdf.set_x(22)
                pdf.set_font('Helvetica', '', 7.5)
                pdf.set_text_color(40, 40, 40)
                pdf.write(3.6, f"* {c_desc}: ")
                pdf.set_font('Helvetica', 'B', 7.5)
                if is_crit_full:
                    pdf.set_text_color(34, 139, 34)
                    pdf.write(3.6, f"[Awarded {c_awarded_disp} / {c_max_disp} Mark{'s' if c_max_disp > 1 else ''}]\n")
                elif c_awarded > 0:
                    pdf.set_text_color(217, 119, 6)
                    pdf.write(3.6, f"[Partial {c_awarded_disp} / {c_max_disp} Marks]\n")
                else:
                    pdf.set_text_color(220, 38, 38)
                    pdf.write(3.6, f"[0 / {c_max_disp} Marks]\n")

        # Question Total Awarded Badge
        is_full = (isinstance(awarded, (int, float)) and awarded >= max_m)
        is_zero = (awarded == 0)
        verdict_tag = "FULL CREDIT" if is_full else ("NIL CREDIT" if is_zero else "PARTIAL CREDIT")
        mark_summary = f">> Question {q_num} Total Awarded: {awarded} / {max_m} Marks [{verdict_tag}]"

        pdf.set_x(18)
        pdf.set_font('Helvetica', 'B', 8)
        pdf.set_text_color(220, 38, 38)
        pdf.cell(178, 4, mark_summary, ln=1)

        if reasoning:
            pdf.set_x(22)
            pdf.set_font('Helvetica', 'I', 7.5)
            pdf.set_text_color(185, 28, 28)
            pdf.multi_cell(174, 3.6, f"Assessor Note: {reasoning}")

        pdf.ln(2.5)

    # ── Section A ─────────────────────────────────────────────────────────────
    if pdf.get_y() > 225:
        pdf.add_page()
    else:
        pdf.set_draw_color(0, 9, 83)
        pdf.set_line_width(0.4)
        pdf.line(14, pdf.get_y(), 196, pdf.get_y())
        pdf.ln(2.5)

    pdf.set_font('Helvetica', 'B', 10)
    pdf.set_text_color(0, 9, 83)
    pdf.cell(130, 5, "SECTION A (30 MARKS) -- Answer ALL Questions", ln=0)
    pdf.set_font('Helvetica', 'B', 8.5)
    pdf.set_text_color(220, 38, 38)
    pdf.cell(52, 5, f"Subtotal: {sec_a_subtotal} / 30 Marks", ln=1, align='R')
    pdf.ln(1.5)

    for idx, item in enumerate(sec_a_items):
        render_marked_question(item, idx + 1, is_section_b=False)

    # ── Section B ─────────────────────────────────────────────────────────────
    if pdf.get_y() > 215:
        pdf.add_page()
    else:
        pdf.set_draw_color(0, 9, 83)
        pdf.set_line_width(0.4)
        pdf.line(14, pdf.get_y(), 196, pdf.get_y())
        pdf.ln(2.5)

    pdf.set_font('Helvetica', 'B', 10)
    pdf.set_text_color(0, 9, 83)
    pdf.cell(130, 5, "SECTION B (40 MARKS) -- Answer ANY TWO Questions", ln=0)
    pdf.set_font('Helvetica', 'B', 8.5)
    pdf.set_text_color(220, 38, 38)
    pdf.cell(52, 5, f"Subtotal: {sec_b_subtotal} / 40 Marks", ln=1, align='R')
    pdf.ln(1.5)

    for idx, item in enumerate(sec_b_items):
        render_marked_question(item, idx + 11, is_section_b=True)

    # ── Overall Feedback & Signatures Footer ──────────────────────────────────
    if pdf.get_y() > 220:
        pdf.add_page()
    else:
        pdf.set_draw_color(0, 9, 83)
        pdf.set_line_width(0.4)
        pdf.line(14, pdf.get_y(), 196, pdf.get_y())
        pdf.ln(3)

    pdf.set_font('Helvetica', 'B', 8.5)
    pdf.set_text_color(0, 9, 83)
    pdf.cell(0, 4.5, "ASSESSOR / TRAINER EXAMINATION FEEDBACK REMARKS:", ln=1)

    comments = clean_pdf_text(payload.get("trainer_comments") or "The candidate demonstrates exceptional competence in digital literacy principles, practical workplace computer procedures, and software applications.")
    rem_y = pdf.get_y()
    pdf.set_fill_color(255, 251, 235)
    pdf.set_draw_color(217, 119, 6)
    pdf.set_line_width(0.3)
    pdf.rect(14, rem_y, 182, 12, 'DF')
    pdf.set_xy(16, rem_y + 1.5)
    pdf.set_font('Helvetica', 'I', 8)
    pdf.set_text_color(60, 40, 0)
    pdf.multi_cell(178, 3.8, f'"{comments}"')

    pdf.set_y(rem_y + 14)
    sig_y = pdf.get_y()
    if sig_y > 245:
        pdf.add_page()
        sig_y = pdf.get_y()

    pdf.set_fill_color(248, 250, 252)
    pdf.set_draw_color(180, 190, 205)
    pdf.rect(14, sig_y, 182, 22, 'DF')

    pdf.set_xy(16, sig_y + 1.5)
    pdf.set_font('Helvetica', 'B', 8)
    pdf.set_text_color(0, 9, 83)
    pdf.cell(90, 4, f"Assessor / Trainer: {trainer_n}", ln=0)
    pdf.cell(90, 4, "Internal Verifier / HOD: Computing & Informatics", ln=1)

    pdf.set_xy(16, sig_y + 6.5)
    pdf.set_font('Helvetica', '', 7.5)
    pdf.set_text_color(40, 40, 40)
    pdf.cell(90, 4, f"Signature: {trainer_n} (Digitally Certified Assessor)", ln=0)
    pdf.cell(90, 4, "Signature: ____________________________________", ln=1)

    pdf.set_xy(16, sig_y + 11.5)
    pdf.cell(90, 4, f"Date: {eval_date}", ln=0)
    pdf.cell(90, 4, "Official Stamp: Mukiria Assessment Center 01200004", ln=1)

    out = pdf.output(dest='S')
    return out.encode('latin1') if isinstance(out, str) else bytes(out)


def generate_timetable_pdf(payload):
    pdf = FPDF(orientation='L', unit='mm', format='A4')
    pdf.set_margins(8, 8, 8)
    pdf.set_auto_page_break(False)
    pdf.add_page(orientation='L')

    current_dir = os.path.dirname(os.path.abspath(__file__))
    project_root = os.path.dirname(current_dir)
    logo_path = os.path.join(project_root, "client", "public", "mtti-logo.jpg")
    if not os.path.exists(logo_path):
        logo_path = os.path.abspath("client/public/mtti-logo.jpg")

    # Outer border (single page A4 landscape is 297mm x 210mm)
    pdf.set_draw_color(0, 9, 83)
    pdf.set_line_width(0.7)
    pdf.rect(6, 6, 285, 198)
    pdf.set_draw_color(196, 136, 32)
    pdf.set_line_width(0.3)
    pdf.rect(7.5, 7.5, 282, 195)

    # Header section
    y_hdr = 10
    if os.path.exists(logo_path):
        try:
            pdf.image(logo_path, x=10, y=y_hdr, w=18)
        except Exception:
            pass

    pdf.set_xy(31, y_hdr)
    pdf.set_font('Helvetica', 'B', 15)
    pdf.set_text_color(0, 9, 83)
    pdf.cell(175, 5.5, 'MUKIRIA TECHNICAL TRAINING INSTITUTE', ln=0)

    pdf.set_font('Helvetica', 'B', 8.5)
    pdf.set_text_color(0, 0, 0)
    pdf.cell(78, 5.5, 'Assessment Center: 01200004', ln=1, align='R')

    pdf.set_xy(31, y_hdr + 5.5)
    pdf.set_font('Helvetica', 'B', 10.5)
    pdf.set_text_color(40, 40, 40)
    title_sub = clean_pdf_text(payload.get('schedule_title', 'MASTER ACADEMIC TIMETABLE'))
    pdf.cell(175, 5, title_sub, ln=0)

    pdf.set_font('Helvetica', '', 8)
    pdf.set_text_color(80, 80, 80)
    pdf.cell(78, 5, 'All lectures strictly 2-hour duration', ln=1, align='R')

    pdf.set_xy(31, y_hdr + 11)
    pdf.set_font('Helvetica', '', 8)
    term = payload.get('term', 'Term 3, 2026')
    week = payload.get('week', 6)
    term_meta = clean_pdf_text(f"{term} | Week {week} | Effective: 31 Aug 2026 - 20 Nov 2026")
    pdf.cell(175, 4.5, term_meta, ln=0)

    pdf.set_font('Helvetica', 'B', 8)
    pdf.set_text_color(34, 139, 34)
    pdf.cell(78, 4.5, 'TVET CDACC CERTIFIED', ln=1, align='R')

    # Table Geometry
    # Widths: Day: 18, P1: 60, Tea: 12, P2: 60, Lunch: 12, P3: 60, Brk: 9, P4: 50 -> Sum = 281mm
    col_widths = [18, 60, 12, 60, 12, 60, 9, 50]
    headers = [
        ('Day', ''),
        ('Period 1', '08:00 - 10:00'),
        ('TEA', '10:00'),
        ('Period 2', '10:30 - 12:30'),
        ('LUNCH', '12:30'),
        ('Period 3', '13:30 - 15:30'),
        ('BRK', '15:30'),
        ('Period 4', '15:35 - 17:35')
    ]

    table_x = 8
    table_y = 28
    header_h = 10
    row_h = 27.5 # 5 rows * 27.5 = 137.5mm + 10mm = 147.5mm

    # Render Header Row
    pdf.set_xy(table_x, table_y)
    cur_x = table_x
    for idx, ((h1, h2), w) in enumerate(zip(headers, col_widths)):
        is_break = idx in [2, 4, 6]
        pdf.set_fill_color(240, 243, 248) if not is_break else pdf.set_fill_color(254, 243, 199)
        pdf.set_draw_color(0, 0, 0)
        pdf.set_line_width(0.3)
        pdf.rect(cur_x, table_y, w, header_h, 'DF')

        pdf.set_xy(cur_x, table_y + 1)
        pdf.set_font('Helvetica', 'B', 8.5 if not is_break else 7.5)
        pdf.set_text_color(0, 9, 83) if not is_break else pdf.set_text_color(146, 64, 14)
        pdf.cell(w, 4, h1, align='C', ln=1)

        if h2:
            pdf.set_xy(cur_x, table_y + 5)
            pdf.set_font('Helvetica', '', 7)
            pdf.set_text_color(80, 80, 80)
            pdf.cell(w, 3.5, h2, align='C')

        cur_x += w

    days = ['Monday', 'Tuesday', 'Wednesday', 'Thursday', 'Friday']
    schedule_data = payload.get('schedule', {})

    cur_y = table_y + header_h
    for day_idx, day in enumerate(days):
        day_sessions = schedule_data.get(day, [])
        p1 = next((s for s in day_sessions if s.get('period') == 1), None)
        p2 = next((s for s in day_sessions if s.get('period') == 2), None)
        p3 = next((s for s in day_sessions if s.get('period') == 3), None)
        p4 = next((s for s in day_sessions if s.get('period') == 4), None)

        cur_x = table_x

        # Day cell
        pdf.set_fill_color(248, 250, 252)
        pdf.rect(cur_x, cur_y, col_widths[0], row_h, 'DF')
        pdf.set_xy(cur_x, cur_y + (row_h / 2) - 3)
        pdf.set_font('Helvetica', 'B', 10)
        pdf.set_text_color(0, 9, 83)
        pdf.cell(col_widths[0], 6, day[:3].upper(), align='C')
        cur_x += col_widths[0]

        # Function to render period cell
        def draw_period_cell(x, y, w, h, sess):
            pdf.set_fill_color(255, 255, 255)
            pdf.set_draw_color(0, 0, 0)
            pdf.rect(x, y, w, h, 'DF')
            if not sess:
                pdf.set_xy(x, y + (h / 2) - 2)
                pdf.set_font('Helvetica', '', 9)
                pdf.set_text_color(180, 180, 180)
                pdf.cell(w, 4, '-', align='C')
                return

            pdf.set_xy(x + 1.5, y + 1.5)
            # Unit title
            pdf.set_font('Helvetica', 'B', 8)
            pdf.set_text_color(10, 10, 10)
            unit_t = clean_pdf_text(sess.get('unitTitle', ''))[:38]
            pdf.cell(w - 3, 4, unit_t, ln=1)

            # Class code (cleaned)
            cohort = clean_pdf_text(sess.get('classCode', ''))
            if unit_t and cohort.upper().startswith(unit_t.upper()):
                cohort = cohort[len(unit_t):].strip()
            cohort = re.sub(r'^(?:ICT\s*SKIL[LS]\s*(?:1\s*)?)', '', cohort, flags=re.IGNORECASE).strip()

            pdf.set_x(x + 1.5)
            pdf.set_font('Helvetica', 'B', 7.5)
            pdf.set_text_color(60, 60, 60)
            pdf.cell(w - 3, 3.8, cohort[:35], ln=1)

            # Trainer name
            trainer_n = clean_pdf_text(sess.get('trainerName') or sess.get('trainer') or '')
            if trainer_n:
                pdf.set_x(x + 1.5)
                pdf.set_font('Helvetica', 'I', 7.5)
                pdf.set_text_color(0, 40, 120)
                pdf.cell(w - 3, 3.5, trainer_n[:35], ln=1)

            # Venue pill
            venue_str = clean_pdf_text(sess.get('venue', ''))
            if venue_str:
                pdf.set_xy(x + 1.5, y + h - 5.5)
                pdf.set_fill_color(240, 240, 240)
                pdf.set_draw_color(160, 160, 160)
                pdf.rect(x + 1.5, y + h - 5.5, 22, 4, 'DF')
                pdf.set_xy(x + 1.5, y + h - 5.5)
                pdf.set_font('Helvetica', 'B', 7)
                pdf.set_text_color(180, 20, 20)
                pdf.cell(22, 4, f'RM: {venue_str}', align='C')

        # Period 1
        draw_period_cell(cur_x, cur_y, col_widths[1], row_h, p1)
        cur_x += col_widths[1]

        # Tea break
        pdf.set_fill_color(254, 243, 199)
        pdf.rect(cur_x, cur_y, col_widths[2], row_h, 'DF')
        pdf.set_xy(cur_x, cur_y + (row_h / 2) - 3)
        pdf.set_font('Helvetica', 'B', 6.5)
        pdf.set_text_color(146, 64, 14)
        pdf.cell(col_widths[2], 6, 'TEA', align='C')
        cur_x += col_widths[2]

        # Period 2
        draw_period_cell(cur_x, cur_y, col_widths[3], row_h, p2)
        cur_x += col_widths[3]

        # Lunch
        pdf.set_fill_color(254, 243, 199)
        pdf.rect(cur_x, cur_y, col_widths[4], row_h, 'DF')
        pdf.set_xy(cur_x, cur_y + (row_h / 2) - 3)
        pdf.set_font('Helvetica', 'B', 6.5)
        pdf.set_text_color(146, 64, 14)
        pdf.cell(col_widths[4], 6, 'LUNCH', align='C')
        cur_x += col_widths[4]

        # Period 3
        draw_period_cell(cur_x, cur_y, col_widths[5], row_h, p3)
        cur_x += col_widths[5]

        # Short break
        pdf.set_fill_color(254, 243, 199)
        pdf.rect(cur_x, cur_y, col_widths[6], row_h, 'DF')
        pdf.set_xy(cur_x, cur_y + (row_h / 2) - 2)
        pdf.set_font('Helvetica', '', 6)
        pdf.set_text_color(146, 64, 14)
        pdf.cell(col_widths[6], 4, '-', align='C')
        cur_x += col_widths[6]

        # Period 4
        draw_period_cell(cur_x, cur_y, col_widths[7], row_h, p4)
        cur_x += col_widths[7]

        cur_y += row_h

    # Official Endorsement Footer
    y_ft = cur_y + 3.5
    pdf.set_xy(table_x, y_ft)
    pdf.set_font('Helvetica', 'B', 7.5)
    pdf.set_text_color(0, 0, 0)
    pdf.cell(90, 3.5, 'Prepared by: Timetable Committee & Deputy Principal Academics', ln=0)
    pdf.cell(100, 3.5, 'Academic Registrar Verification: _________________________', ln=0)
    pdf.cell(91, 3.5, 'Official Institutional Stamp', ln=1, align='R')

    pdf.set_xy(table_x, y_ft + 3.5)
    pdf.set_font('Helvetica', '', 7)
    pdf.set_text_color(90, 90, 90)
    pdf.cell(90, 3.5, 'Software: aSc Timetables Institutional Edition 2026', ln=0)
    pdf.cell(100, 3.5, 'Date: October 2026', ln=0)
    pdf.cell(91, 3.5, 'Mukiria TTI Quality Assurance Directorate', ln=1, align='R')

    out = pdf.output(dest='S')
    return out.encode('latin1') if isinstance(out, str) else bytes(out)


class APIHandler(BaseHTTPRequestHandler):
    def send_cors_headers(self):
        self.send_header("Access-Control-Allow-Origin", "*")
        self.send_header("Access-Control-Allow-Headers", "Content-Type,Authorization")
        self.send_header("Access-Control-Allow-Methods", "GET,POST,OPTIONS")

    def do_OPTIONS(self):
        self.send_response(200)
        self.send_cors_headers()
        self.end_headers()

    def do_GET(self):
        try:
            parsed_url = urllib.parse.urlparse(self.path)
            path = parsed_url.path
            query_params = urllib.parse.parse_qs(parsed_url.query)
            
            if path == "/api/health":
                self.send_response(200)
                self.send_cors_headers()
                self.send_header("Content-Type", "application/json")
                self.end_headers()
                self.wfile.write(json.dumps({"status": "ok", "app": "MTTI Trainer System"}).encode("utf-8"))
                return
                
            elif path == "/api/curriculum/departments":
                depts = curriculum_service.get_departments()
                self.send_response(200)
                self.send_cors_headers()
                self.send_header("Content-Type", "application/json")
                self.end_headers()
                self.wfile.write(json.dumps({"success": True, "departments": depts}).encode("utf-8"))
                return
                
            elif path == "/api/curriculum/units":
                dept = query_params.get("department", [None])[0]
                level = query_params.get("level", [None])[0]
                search = query_params.get("search", [None])[0]
                units = curriculum_service.get_units(department=dept, level=level, search=search)
                self.send_response(200)
                self.send_cors_headers()
                self.send_header("Content-Type", "application/json")
                self.end_headers()
                self.wfile.write(json.dumps({"success": True, "count": len(units), "units": units}).encode("utf-8"))
                return
                
            elif path.startswith("/api/curriculum/unit/"):
                raw_code = path.replace("/api/curriculum/unit/", "").strip()
                unit = curriculum_service.get_unit_by_code(raw_code)
                if unit:
                    self.send_response(200)
                    self.send_cors_headers()
                    self.send_header("Content-Type", "application/json")
                    self.end_headers()
                    self.wfile.write(json.dumps({"success": True, "unit": unit}).encode("utf-8"))
                else:
                    self.send_response(404)
                    self.send_cors_headers()
                    self.send_header("Content-Type", "application/json")
                    self.end_headers()
                    self.wfile.write(json.dumps({"error": "Unit not found", "code": raw_code}).encode("utf-8"))
                return
                
            else:
                self.send_response(404)
                self.send_cors_headers()
                self.send_header("Content-Type", "application/json")
                self.end_headers()
                self.wfile.write(json.dumps({"error": "Endpoint not found"}).encode("utf-8"))
                return
        except Exception as e:
            self.send_response(500)
            self.send_cors_headers()
            self.send_header("Content-Type", "application/json")
            self.end_headers()
            self.wfile.write(json.dumps({"error": str(e)}).encode("utf-8"))

    def do_POST(self):
        content_length = int(self.headers.get("Content-Length", 0))
        post_data = self.rfile.read(content_length)
        
        # Determine path
        path = self.path.split("?")[0]
        
        try:
            if path == "/api/curriculum/rescan":
                res = curriculum_service.trigger_rescan()
                self.send_response(200 if res.get("success") else 500)
                self.send_cors_headers()
                self.send_header("Content-Type", "application/json")
                self.end_headers()
                self.wfile.write(json.dumps(res).encode("utf-8"))
                return

            if path == "/api/parse-doc":
                payload = json.loads(post_data.decode("utf-8"))
                file_base64 = payload.get("file_data")
                context = payload.get("context", "learning_plan")
                role = payload.get("role", "trainer")
                
                if not file_base64:
                    self.send_response(400)
                    self.send_cors_headers()
                    self.send_header("Content-Type", "application/json")
                    self.end_headers()
                    self.wfile.write(json.dumps({"error": "No file data provided in upload"}).encode("utf-8"))
                    return

                # Decode file bytes
                file_bytes = base64.b64decode(file_base64)
                
                # If this is a trainee uploading an assignment, bypass the CDACC curriculum parsing
                if context == "trainee_assignment":
                    self.send_response(200)
                    self.send_cors_headers()
                    self.send_header("Content-Type", "application/json")
                    self.end_headers()
                    self.wfile.write(json.dumps({
                        "message": "Assignment uploaded successfully",
                        "status": "success",
                        "grade_status": "pending"
                    }).encode("utf-8"))
                    return

                # Determine file type
                filename = payload.get("filename", "").lower()
                
                full_text = []
                if filename.endswith(".pdf") or payload.get("content_type") == "application/pdf":
                    pdf_extracted = False
                    try:
                        import pdfplumber
                        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
                            for page in pdf.pages:
                                extracted = page.extract_text()
                                if extracted:
                                    full_text.append(extracted)
                                # Extract tables
                                tables = page.extract_tables()
                                for table in tables:
                                    for row in table:
                                        row_cells = [str(cell).strip() for cell in row if cell]
                                        if row_cells:
                                            full_text.append(" | ".join(row_cells))
                            pdf_extracted = True
                    except Exception as e_plumber:
                        print("pdfplumber parsing warning:", e_plumber)

                    if not pdf_extracted:
                        try:
                            import pypdf
                            reader = pypdf.PdfReader(io.BytesIO(file_bytes))
                            for page in reader.pages:
                                txt = page.extract_text()
                                if txt:
                                    full_text.append(txt)
                        except Exception as e_pypdf:
                            print("pypdf fallback error:", e_pypdf)
                else:
                    # Default to docx
                    doc = docx.Document(io.BytesIO(file_bytes))
                    for p in doc.paragraphs:
                        if p.text.strip():
                            full_text.append(p.text.strip())
                            
                    for table in doc.tables:
                        for row in table.rows:
                            row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                            if row_cells:
                                full_text.append(" | ".join(row_cells))
                                
                combined_text = "\n".join(full_text)
                
                # Resilient CDACC Unit Code finder
                unit_code = None
                course_name = ""
                unit_name = ""
                
                # Extract Course Name and Unit Name
                course_name_match = re.search(r'(?:Course Name|Course|Programme):\s*(.*)', combined_text, re.IGNORECASE)
                if course_name_match:
                    course_name = course_name_match.group(1).strip()
                    
                unit_name_match = re.search(r'(?:Unit Name|Module Name|Subject):\s*(.*)', combined_text, re.IGNORECASE)
                if unit_name_match:
                    unit_name = unit_name_match.group(1).strip()
                
                # 1. Look for explicit unit keyword
                unit_match = re.search(r'(?:Unit Code|Unit Code of Competency|Module Code|Code):\s*([A-Z0-9/\-]+)', combined_text, re.IGNORECASE)
                if unit_match:
                    unit_code = unit_match.group(1).strip()
                else:
                    # 2. Search standard CDACCSurveys formats: (e.g. CON/OS/SUR/CR/01/6/A or SUR/OS/SUR/CR/01/6)
                    cdacc_match = re.search(r'\b([A-Z0-9]{2,6}/OS/[A-Z0-9]{2,6}/[A-Z0-9]{2,6}/\d{2}/\d(?:/[A-Z0-9])?)\b', combined_text, re.IGNORECASE)
                    if cdacc_match:
                        unit_code = cdacc_match.group(1).strip()
                    else:
                        # 3. Match any slash separated string containing level digit
                        general_match = re.search(r'\b([A-Z0-9]{2,6}/[A-Z0-9]{2,6}/[A-Z0-9]{2,6}/\d)\b', combined_text, re.IGNORECASE)
                        if general_match:
                            unit_code = general_match.group(1).strip()
                
                if not unit_code:
                    unit_code = "CON/OS/SUR/CR/01/6"  # Default to land survey if failed to parse

                # Trainer & Class & Series extraction
                trainer_match = re.search(r'(?:Trainer|Name|Instructor|Author):\s*(.*)', combined_text, re.IGNORECASE)
                class_match = re.search(r'(?:Class Code|Class):\s*(.*)', combined_text, re.IGNORECASE)
                series_match = re.search(r'(?:Series|Term|Semester):\s*(.*)', combined_text, re.IGNORECASE)
                
                trainer = trainer_match.group(1).strip() if trainer_match else "MR. ALEXANDER KINOTI"
                class_code = class_match.group(1).strip() if class_match else ""
                series = series_match.group(1).strip() if series_match else ""
                
                # Parse outcomes & topics
                topic_match = re.findall(r'(?:Topic|Lesson|Session Title|Module):\s*(.*)', combined_text, re.IGNORECASE)
                outcomes_match = re.findall(r'(?:Outcome|Objective|Trainee should be able to|LO|Element):\s*(.*)', combined_text, re.IGNORECASE)
                
                topics = [t.strip() for t in topic_match if t.strip()]
                raw_outcomes = [o.strip() for o in outcomes_match if o.strip()]
                
                # Filter out table header rows and garbage from outcomes
                def clean_outcome_text(text):
                    t = text.strip()
                    low = t.lower()
                    if "suggested assessment" in low or "content |" in low or "| suggested" in low or "suggested assessmentmethods" in low:
                        return None
                    if low in ["learning outcome", "learning outcomes", "content", "topic", "session title", "objectives"]:
                        return None
                    return t

                valid_outcomes = [clean_outcome_text(o) for o in raw_outcomes if clean_outcome_text(o)]
                
                # Check for Land Surveying, Civil, or other specific domain keywords if unit code is unassigned or generic
                combined_low = combined_text.lower()
                department = "Computing & Informatics"
                if any(k in combined_low for k in ["geodetic", "cadastre", "cadastral", "topograph", "land survey", "traverse", "beacon", "theodolite", "tacheometry", "levelling"]):
                    department = "Land Survey and Geomatics"
                    if not unit_code or len(unit_code) < 6 or "/" not in unit_code:
                        unit_code = "CON/OS/SUR/CR/01/6"
                        unit_name = "Conduct Geodetic and Cadastral Surveys"
                        course_name = "Diploma in Land Survey & Geomatics (Level 6)"
                        class_code = class_code or "LS6/M/S/24"
                elif any(k in combined_low for k in ["civil", "concrete", "structural", "masonry", "carpentry", "plumbing", "building drawing"]):
                    department = "Civil and Building Engineering"
                    if not unit_code or len(unit_code) < 6 or "/" not in unit_code:
                        unit_code = "CON/OS/CM/CR/01/6"
                        unit_name = "Apply Construction Management Principles"
                        course_name = "Diploma in Civil Engineering (Level 6)"
                        class_code = class_code or "CE6/M/S/24"
                elif any(k in combined_low for k in ["business", "office administration", "accounting", "secretarial", "human resource"]):
                    department = "Business and Office Administration"
                    if not unit_code or len(unit_code) < 6 or "/" not in unit_code:
                        unit_code = "BUS/OS/BM/CR/01/6"
                        unit_name = "Office Management and Administration"
                        course_name = "Diploma in Business Management (Level 6)"
                        class_code = class_code or "BUS6/M/25"

                level = get_level_from_code(unit_code)
                
                # Cross-reference with institutional curriculum database (from D:\Curriculum and OS)
                matched_db_unit = curriculum_service.get_unit_by_code(unit_code)
                if not matched_db_unit and unit_name:
                    matched_db_unit = curriculum_service.get_unit_by_code(unit_name)
                    
                elements = []
                weeks_breakdown = []
                if matched_db_unit:
                    unit_code = matched_db_unit.get("unit_code", unit_code)
                    if not unit_name or len(unit_name) < 4:
                        unit_name = matched_db_unit.get("unit_title", unit_name)
                    if not course_name or len(course_name) < 4:
                        course_name = matched_db_unit.get("course", course_name)
                    if matched_db_unit.get("department"):
                        department = matched_db_unit.get("department")
                    level = matched_db_unit.get("level", level)
                    elements = matched_db_unit.get("elements", [])
                    weeks_breakdown = matched_db_unit.get("weeks_breakdown", [])
                    db_outcomes = [lo["title"] for lo in matched_db_unit.get("learning_outcomes", [])]
                    if db_outcomes:
                        valid_outcomes = db_outcomes
                    db_topics = []
                    for lo in matched_db_unit.get("learning_outcomes", []):
                        db_topics.extend(lo.get("content", []))
                    if db_topics:
                        topics = db_topics[:12]

                # Ensure outcomes are concrete and informative
                if not valid_outcomes or len(valid_outcomes) < len(topics):
                    enhanced_outcomes = []
                    for t in topics:
                        enhanced_outcomes.append(f"Demonstrate theoretical understanding and practical competency in {t} adhering to CDACC guidelines.")
                    outcomes = enhanced_outcomes
                else:
                    outcomes = valid_outcomes

                # Warnings
                warnings = []

                response_data = {
                    "success": True,
                    "trainer": trainer,
                    "unit_code": unit_code,
                    "course_name": course_name,
                    "unit_name": unit_name,
                    "department": department,
                    "class_code": class_code,
                    "series": series or "SEP - NOV 2026",
                    "level": level,
                    "topics": topics,
                    "outcomes": outcomes,
                    "elements": elements,
                    "weeks_breakdown": weeks_breakdown,
                    "warnings": warnings
                }
                
                self.send_response(200)
                self.send_cors_headers()
                self.send_header("Content-Type", "application/json")
                self.end_headers()
                self.wfile.write(json.dumps(response_data).encode("utf-8"))
                
            elif path == "/api/generate-exam":
                payload = json.loads(post_data.decode("utf-8"))
                unit_code = payload.get("unit_code", "ICT/OS/CS/CR/11/6/A")
                course_name = payload.get("course_name", "Software Engineering Fundamentals")
                level = payload.get("level")
                
                if level is None:
                    level = get_level_from_code(unit_code)
                    
                topics = payload.get("topics", [])
                outcomes = payload.get("outcomes", [])
                is_practical = payload.get("is_practical", False)
                
                if is_practical:
                    exam_payload = build_cdacc_practical(unit_code, course_name, level, topics, outcomes)
                else:
                    exam_payload = build_cdacc_exam(unit_code, course_name, level, topics, outcomes)
                
                response_data = {
                    "success": True,
                    "unit_code": unit_code,
                    "course_name": course_name,
                    "payload": exam_payload
                }
                
                self.send_response(200)
                self.send_cors_headers()
                self.send_header("Content-Type", "application/json")
                self.end_headers()
                self.wfile.write(json.dumps(response_data).encode("utf-8"))

            elif path == "/api/complete-session":
                user_role = self.headers.get("x-user-role", "").strip().lower()
                if user_role == "trainee":
                    self.send_response(403)
                    self.send_cors_headers()
                    self.send_header("Content-Type", "application/json")
                    self.end_headers()
                    self.wfile.write(json.dumps({"error": "Forbidden: Trainees cannot complete session plans"}).encode("utf-8"))
                    return

                payload = json.loads(post_data.decode("utf-8"))
                session_plan_id = payload.get("session_plan_id")
                unit_code = payload.get("unit_code", "COMP-204")
                week_number = payload.get("week_number", 1)
                
                row_draft = {
                    "session_plan_id": session_plan_id,
                    "date_delivered": "2026-07-07",
                    "work_actually_covered": f"Delivered learning outcomes matching week {week_number} syllabus plan.",
                    "trainees_present": 42,
                    "signature": "Alexander Kinoti",
                    "status": "unsigned"
                }
                
                formative_questions = [
                    {
                        "id": "fq-1",
                        "text": f"Define the core operating workflow configured for week {week_number} units.",
                        "type": "short_answer",
                        "marks": 5
                    },
                    {
                        "id": "fq-2",
                        "text": f"Explain why redundancy elimination is vital during database schema creation.",
                        "type": "short_answer",
                        "marks": 5
                    }
                ]
                
                response_data = {
                    "success": True,
                    "record_of_work_draft": row_draft,
                    "formative_assessment_draft": {
                        "title": f"Formative Quiz — Week {week_number} ({unit_code})",
                        "questions": formative_questions
                    },
                    "timestamp": "2026-07-07T23:55:00Z"
                }
                
                self.send_response(200)
                self.send_cors_headers()
                self.send_header("Content-Type", "application/json")
                self.end_headers()
                self.wfile.write(json.dumps(response_data).encode("utf-8"))
            elif path == "/api/auto-grade":
                user_role = self.headers.get("x-user-role", "").strip().lower()
                payload = json.loads(post_data.decode("utf-8"))
                student_answer = str(payload.get("student_answer", "")).strip()
                correct_answer = str(payload.get("correct_answer", "")).strip()
                max_marks = float(payload.get("marks", 0))
                question_type = str(payload.get("question_type", "short_answer")).lower()
                regex_pattern = payload.get("regex_pattern")
                keywords = payload.get("keywords", [])
                eval_mode = str(payload.get("evaluation_mode", "semi_objective")).lower()
                requires_review = bool(payload.get("requires_trainer_review", False))

                auto_score = 0.0
                reasoning = ""
                confidence = 1.0
                flagged = False
                rule_applied = "heuristic"

                if not student_answer:
                    auto_score = 0.0
                    reasoning = "No response provided by candidate."
                    confidence = 1.0
                    flagged = False
                    rule_applied = "empty_response"

                # 1. Deterministic MCQ / True-False Exact Option Check
                elif question_type in ["mcq", "true_false"] or eval_mode == "objective":
                    is_correct = (student_answer.lower() == correct_answer.lower())
                    auto_score = max_marks if is_correct else 0.0
                    if is_correct:
                        reasoning = "Correct objective answer selected."
                    else:
                        # Shield answer key if caller is trainee
                        reasoning = "Incorrect choice selected." if user_role == "trainee" else f"Incorrect choice. Expected option: {correct_answer}."
                    confidence = 1.0
                    flagged = False
                    rule_applied = "exact_match"

                # 2. Regex Pattern Matching (Keyboard shortcuts, Excel formulas, exact syntax)
                elif regex_pattern and re.search(regex_pattern, student_answer, re.IGNORECASE):
                    auto_score = max_marks
                    reasoning = f"Response matches required syntax and technical pattern."
                    confidence = 0.95
                    flagged = False
                    rule_applied = "regex_match"

                # 3. Subjective Practical Tasks & Applied Scenarios (Flagged for Review)
                elif question_type in ["practical", "essay"] or eval_mode == "subjective" or requires_review:
                    def tokenize(text):
                        text = re.sub(r'[^\w\s]', '', text.lower())
                        return set(text.split())

                    student_tokens = tokenize(student_answer)
                    kw_found = [kw for kw in keywords if kw.lower() in student_answer.lower()] if keywords else []
                    kw_ratio = (len(kw_found) / len(keywords)) if keywords else 0.5
                    
                    # Award baseline between 40% and 85% based on length and technical keywords
                    length_factor = min(len(student_tokens) / 30.0, 1.0)
                    baseline_ratio = (kw_ratio * 0.6) + (length_factor * 0.4)
                    auto_score = round(max_marks * max(0.4, min(baseline_ratio, 0.85)), 1)
                    confidence = 0.65
                    flagged = True
                    rule_applied = "subjective_review_flagged"
                    reasoning = f"Applied practical task: initial baseline awarded ({auto_score}/{max_marks}). Flagged for trainer review."

                # 4. Semi-Objective Short Answers (Keyword recall + semantic token overlap)
                else:
                    def tokenize(text):
                        text = re.sub(r'[^\w\s]', '', text.lower())
                        return set(text.split())

                    student_tokens = tokenize(student_answer)
                    correct_tokens = tokenize(correct_answer)

                    if not correct_tokens:
                        auto_score = round(max_marks * 0.5, 1)
                        reasoning = "No reference rubric provided. Defaulted to 50% for manual review."
                        confidence = 0.4
                        flagged = True
                        rule_applied = "no_rubric_fallback"
                    else:
                        overlap = student_tokens.intersection(correct_tokens)
                        recall = len(overlap) / len(correct_tokens) if correct_tokens else 0
                        seq_match = difflib.SequenceMatcher(None, student_answer.lower(), correct_answer.lower()).ratio()

                        kw_found = [kw for kw in keywords if kw.lower() in student_answer.lower()] if keywords else []
                        kw_ratio = (len(kw_found) / len(keywords)) if keywords else recall

                        final_ratio = (kw_ratio * 0.45) + (recall * 0.35) + (seq_match * 0.20)

                        if len(student_tokens) > len(correct_tokens) * 0.5 and final_ratio < 0.4:
                            final_ratio += 0.15

                        final_ratio = min(max(final_ratio, 0.0), 1.0)
                        auto_score = round(max_marks * final_ratio, 1)
                        confidence = round(0.6 + (final_ratio * 0.35), 2)

                        if final_ratio >= 0.75:
                            reasoning = f"Strong alignment ({round(final_ratio*100)}%). Core technical concepts demonstrated."
                            flagged = False
                        elif final_ratio >= 0.45:
                            missing = correct_tokens - student_tokens
                            missing_sample = ", ".join(list(missing)[:3]) if missing else "key terms"
                            reasoning = f"Partial alignment ({round(final_ratio*100)}%). Missing some expected concepts like: {missing_sample}."
                            flagged = False
                        else:
                            reasoning = f"Low relevance ({round(final_ratio*100)}%). Flagged for trainer review."
                            flagged = True
                        rule_applied = "keyword_and_semantic_overlap"

                response_data = {
                    "score": auto_score,
                    "max_marks": max_marks,
                    "reasoning": reasoning,
                    "confidence": confidence,
                    "flagged_for_review": flagged,
                    "rule_applied": rule_applied
                }

                self.send_response(200)
                self.send_cors_headers()
                self.send_header("Content-Type", "application/json")
                self.end_headers()
                self.wfile.write(json.dumps(response_data).encode("utf-8"))

            elif path == "/api/export-docx":
                payload = json.loads(post_data.decode("utf-8"))
                unit_code = payload.get("unit_code", "COMP-204")
                course_name = payload.get("course_name", "Software Engineering")
                unit_name = payload.get("unit_name", "")
                class_code = payload.get("class_code", "")
                series = payload.get("series", "")
                exam = payload.get("payload", {})
                level = get_level_from_code(unit_code)
                
                meta_course_name = course_name.upper() if course_name else "ICT TECHNICIAN LEVEL 6"
                meta_course_code = unit_code.upper() if unit_code else "061006T91CT"
                meta_unit_name = unit_name.upper() if unit_name else course_name.upper()
                meta_unit_code = unit_code.upper() if unit_code else "ICT/OS/ENV/CR/01/6"
                meta_class = class_code.upper() if class_code else "ITECH6/M/24"
                meta_series = series if series else "Jan./April 2026"
                
                # Resilient matching for unit specifics (legacy fallback removed to prioritize user inputs)
                
                # Add Brand Logo if exists
                current_dir = os.path.dirname(os.path.abspath(__file__))
                project_root = os.path.dirname(current_dir)
                logo_path = os.path.join(project_root, "client", "public", "mtti-logo.jpg")
                
                doc = docx.Document()
                style = doc.styles['Normal']
                style.font.name = 'Maiandra GD'
                
                if os.path.exists(logo_path):
                    p_logo = doc.add_paragraph()
                    p_logo.alignment = 1 # Center
                    p_logo.add_run().add_picture(logo_path, width=docx.shared.Inches(1.2))
                
                # Center aligned brand headers
                p = doc.add_paragraph()
                p.alignment = 1 # Center
                p.paragraph_format.space_after = docx.shared.Pt(0)
                
                r = p.add_run("MUKIRIA TECHNICAL TRAINING INSTITUTE\n")
                r.bold = True
                r.font.size = docx.shared.Pt(14)
                
                r2 = p.add_run("\nCOMPUTING AND INFORMATICS DEPARTMENT\n")
                r2.bold = True
                r2.font.name = "Maiandra GD"
                r2.font.size = docx.shared.Pt(12)
                
                is_practical = "practical" in exam.get("title", "").lower()
                assessment_title = "WRITTEN ASSESSMENT I"
                if is_practical:
                    assessment_title = "FORMATIVE PRACTICAL ASSESSMENT I"
                elif "ii" in exam.get("title", "").lower() or "2" in exam.get("title", "").lower():
                    assessment_title = "WRITTEN ASSESSMENT II"
                elif "iii" in exam.get("title", "").lower() or "3" in exam.get("title", "").lower():
                    assessment_title = "WRITTEN ASSESSMENT III"
                    
                r3 = p.add_run(f"{assessment_title}\n")
                r3.bold = True
                r3.font.name = "Maiandra GD"
                r3.font.size = docx.shared.Pt(12)
                
                r4 = p.add_run(f"TIME: {exam.get('duration_minutes', 120) // 60} HOURS\n\n")
                r4.bold = True
                r4.font.size = docx.shared.Pt(11)
                
                # Left aligned metadata fields
                p_meta = doc.add_paragraph()
                p_meta.paragraph_format.line_spacing = 1.0
                p_meta.paragraph_format.space_after = docx.shared.Pt(0)
                
                fields = [
                    ("COURSE NAME", meta_course_name),
                    ("COURSE CODE", meta_course_code),
                    ("UNIT NAME", meta_unit_name),
                    ("UNIT CODE", meta_unit_code),
                    ("CLASS", meta_class),
                    ("SERIES", meta_series)
                ]
                
                for key, val in fields:
                    r_key = p_meta.add_run(f"{key}: ")
                    r_key.bold = True
                    p_meta.add_run(f"{val}\n")
                    
                # Instructions to candidate
                p_inst_hdr = doc.add_paragraph()
                p_inst_hdr.paragraph_format.space_before = docx.shared.Pt(12)
                r_inst_hdr = p_inst_hdr.add_run("INSTRUCTIONS TO THE CANDIDATE:")
                r_inst_hdr.bold = True
                r_inst_hdr.font.size = docx.shared.Pt(11)
                
                instructions = [
                    "Write your NAME, ADM, CLASS and DATE on the answer booklet provided",
                    "Answer all questions in section A",
                    f"Answer ANY three questions in section B" if level in [5, 6] else "Answer all questions in Section A and follow instructions in Section B."
                ]
                for inst in instructions:
                    p_inst = doc.add_paragraph(style='List Bullet')
                    p_inst.paragraph_format.space_after = docx.shared.Pt(0)
                    r_inst = p_inst.add_run(inst)
                is_practical = exam.get("type") == "practical" or "practical" in exam.get("title", "").lower()
                
                if is_practical:
                    # Render Practical Assessment Format
                    # Project Brief
                    p_brief = doc.add_paragraph()
                    r_brief_hdr = p_brief.add_run("PROJECT BRIEF\n")
                    r_brief_hdr.bold = True
                    p_brief.add_run(exam.get("project_brief", ""))
                    
                    # Elements Covered
                    p_elements = doc.add_paragraph()
                    r_elem_hdr = p_elements.add_run("Elements Covered:\n")
                    r_elem_hdr.bold = True
                    for elem in exam.get("elements_covered", []):
                        p_elem = doc.add_paragraph(style='List Bullet')
                        p_elem.add_run(elem)
                        
                    # Tasks
                    for task in exam.get("tasks", []):
                        p_task = doc.add_paragraph()
                        r_task_hdr = p_task.add_run(f"{task.get('title')} ({task.get('marks')} Marks)\n")
                        r_task_hdr.bold = True
                        for detail in task.get("details", []):
                            p_detail = doc.add_paragraph(style='List Bullet')
                            p_detail.add_run(detail)
                            
                    # Rubrics
                    for r_idx, rubric_cat in enumerate(exam.get("rubric", [])):
                        p_rubric = doc.add_paragraph()
                        p_rubric.paragraph_format.space_before = docx.shared.Pt(12)
                        
                        table = doc.add_table(rows=1, cols=4)
                        table.style = 'Table Grid'
                        
                        hdr_cells = table.rows[0].cells
                        hdr_cells[0].text = 'No.'
                        hdr_cells[1].text = 'Items to be Evaluated: Kindly award marks as appropriate.'
                        hdr_cells[2].text = 'Max Marks'
                        hdr_cells[3].text = 'Awarded'
                        
                        # Style headers
                        for cell in hdr_cells:
                            for p_cell in cell.paragraphs:
                                p_cell.alignment = 1
                                for run in p_cell.runs:
                                    run.bold = True
                                    run.font.name = "Maiandra GD"
                                    
                        # Add criteria rows
                        for c_idx, criteria in enumerate(rubric_cat.get("criteria", [])):
                            row_cells = table.add_row().cells
                            row_cells[0].text = str(r_idx + 1) + chr(97 + c_idx)
                            row_cells[1].text = criteria.get("description", "")
                            row_cells[2].text = str(criteria.get("marks", 0))
                            row_cells[3].text = ""
                            
                            # Style rows
                            for cell in row_cells:
                                for p_cell in cell.paragraphs:
                                    for run in p_cell.runs:
                                        run.font.name = "Maiandra GD"
                else:
                    # Render Written Assessment Format
                    # Section A
                    sec_a = exam.get("section_a", {})
                    p_sec_a_hdr = doc.add_paragraph()
                    p_sec_a_hdr.paragraph_format.space_before = docx.shared.Pt(12)
                    r_sec_a_hdr = p_sec_a_hdr.add_run(f"SECTION A: ({sec_a.get('total_marks', 40)} MARKS)")
                    r_sec_a_hdr.bold = True
                    r_sec_a_hdr.font.size = docx.shared.Pt(11)
                    
                    for idx, q in enumerate(sec_a.get("questions", [])):
                        p_q = doc.add_paragraph()
                        r_q_num = p_q.add_run(f"{idx+1}. ")
                        r_q_num.bold = True
                        
                        r_q_text = p_q.add_run(f"{q.get('text', '')} ")
                        
                        r_q_marks = p_q.add_run(f"({q.get('marks', 0)} Marks)")
                        r_q_marks.bold = True
                        
                        # Options rendering removed as requested
                                
                    # Section B
                    sec_b = exam.get("section_b", {})
                    p_sec_b_hdr = doc.add_paragraph()
                    p_sec_b_hdr.paragraph_format.space_before = docx.shared.Pt(12)
                    sec_b_title = "SECTION TWO" if level in [5, 6] else "SECTION B"
                    r_sec_b_hdr = p_sec_b_hdr.add_run(f"{sec_b_title} ({sec_b.get('total_marks', 60)} MARKS)")
                    r_sec_b_hdr.bold = True
                    r_sec_b_hdr.font.size = docx.shared.Pt(11)
                    
                    p_sec_b_inst = doc.add_paragraph()
                    r_sec_b_inst = p_sec_b_inst.add_run(sec_b.get("instructions", "Answer subjective items as guided."))
                    r_sec_b_inst.italic = True
                    
                    for idx, q in enumerate(sec_b.get("questions", [])):
                        p_q = doc.add_paragraph()
                        # Number starts after section A questions (11 or 6)
                        sec_a_len = len(sec_a.get("questions", []))
                        r_q_num = p_q.add_run(f"{idx + sec_a_len + 1}. ")
                        r_q_num.bold = True
                        
                        # Note: In main.py build_cdacc_exam, we already formatted a, b, c inside text!
                        # We just need to split it by new lines to render cleanly
                        q_parts = q.get('text', '').split('\n')
                        for part_idx, part in enumerate(q_parts):
                            if part_idx == 0:
                                p_q.add_run(f"{part} ")
                            else:
                                p_sub = doc.add_paragraph()
                                p_sub.paragraph_format.left_indent = docx.shared.Inches(0.5)
                                p_sub.add_run(part)
                                
                        r_q_marks = p_q.add_run(f"({q.get('marks', 0)} Marks)")
                        r_q_marks.bold = True
                
                # Save to BytesIO
                f_stream = io.BytesIO()
                doc.save(f_stream)
                doc_base64 = base64.b64encode(f_stream.getvalue()).decode("utf-8")
                
                response_data = {
                    "success": True,
                    "filename": f"{unit_code.replace('/', '_')}_Exam.docx",
                    "file_data": doc_base64
                }
                
                self.send_response(200)
                self.send_cors_headers()
                self.send_header("Content-Type", "application/json")
                self.end_headers()
                self.wfile.write(json.dumps(response_data).encode("utf-8"))

            elif path == "/api/export-session-plan":
                # ── MTTI/F/CUR/05 Session Plan ─────────────────────────────────
                payload = json.loads(post_data.decode("utf-8"))
                plan = payload.get("payload", {})
                steps = plan.get("delivery_steps", [])
                n_step_rows = max(1, len(steps))

                doc = docx.Document()

                # Set page margins to narrow
                from docx.oxml.ns import qn as _qn
                from docx.oxml import OxmlElement as _OxmlElement
                section = doc.sections[0]
                section.top_margin    = docx.shared.Inches(0.6)
                section.bottom_margin = docx.shared.Inches(0.6)
                section.left_margin   = docx.shared.Inches(0.8)
                section.right_margin  = docx.shared.Inches(0.8)

                # Set default font
                style = doc.styles['Normal']
                style.font.name = 'Maiandra GD'
                style.font.size = docx.shared.Pt(11)

                # ── Header paragraphs ────────────────────────────────────────────
                current_dir  = os.path.dirname(os.path.abspath(__file__))
                project_root = os.path.dirname(current_dir)
                logo_path    = os.path.join(project_root, "client", "public", "mtti-logo.jpg")

                p_code = doc.add_paragraph()
                p_code.alignment = 2  # Right aligned
                p_code.add_run("MTTI/F/CUR/05").bold = True

                p_hdr = doc.add_paragraph()
                p_hdr.alignment = 1  # Center
                if os.path.exists(logo_path):
                    p_hdr.add_run().add_picture(logo_path, width=docx.shared.Inches(1.0))
                r_inst = p_hdr.add_run("\nMUKIRIA TECHNICAL TRAINING INSTITUTE\n\n")
                r_inst.bold = True
                r_inst.font.size = docx.shared.Pt(14)
                
                r_sp = p_hdr.add_run("SESSION PLAN")
                r_sp.bold = True

                total_rows = 15 + n_step_rows + 7
                tbl = doc.add_table(rows=total_rows, cols=4)
                tbl.style = 'Table Grid'
                
                # Set precise 4-column widths matching the image exactly (15%, 35%, 25%, 25%)
                # Total width = 6.9 inches (8.5 - 0.8*2)
                from docx.shared import Inches
                widths = [Inches(1.035), Inches(2.415), Inches(1.725), Inches(1.725)]
                for row in tbl.rows:
                    for idx, width in enumerate(widths):
                        row.cells[idx].width = width

                def _set_labeled_cell(cell, label, value=""):
                    cell.text = ""
                    p = cell.paragraphs[0]
                    run_label = p.add_run(label)
                    run_label.bold = True
                    run_label.font.name = 'Maiandra GD'
                    if value:
                        run_val = p.add_run(" " + str(value))
                        run_val.font.name = 'Maiandra GD'

                def _set_header_cell(cell, text, fill="D9D9D9", align=None):
                    cell.text = ""
                    p = cell.paragraphs[0]
                    if align is not None:
                        p.alignment = align
                    run = p.add_run(text)
                    run.bold = True
                    run.font.name = 'Maiandra GD'
                    if fill:
                        from docx.oxml.ns import nsdecls
                        from docx.oxml import parse_xml
                        tcPr = cell._tc.get_or_add_tcPr()
                        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill}"/>')
                        tcPr.append(shd)

                def _merge(row, start_col, end_col):
                    row.cells[start_col].merge(row.cells[end_col])

                # Row 0
                _set_labeled_cell(tbl.rows[0].cells[0], "Date:", plan.get("date", ""))
                _merge(tbl.rows[0], 0, 1) # Left side (50%)
                _set_labeled_cell(tbl.rows[0].cells[2], "Time:", plan.get("time_duration", ""))
                _merge(tbl.rows[0], 2, 3) # Right side (50%)
                tbl.rows[0].cells[2].merge(tbl.rows[1].cells[3]) # Vertical merge on right

                # Row 1
                _set_labeled_cell(tbl.rows[1].cells[0], "Week:", plan.get("week_number", ""))
                _merge(tbl.rows[1], 0, 1)

                # Row 2
                _set_labeled_cell(tbl.rows[2].cells[0], "Trainer name:", plan.get("trainer_name", "Alexander Kinoti"))
                _merge(tbl.rows[2], 0, 1)
                _set_labeled_cell(tbl.rows[2].cells[2], "Department:", plan.get("department", "Computing & Informatics"))
                _merge(tbl.rows[2], 2, 3)

                # Row 3, 4, 5
                _set_labeled_cell(tbl.rows[3].cells[0], "Unit of Competence:", plan.get("unit_name", ""))
                _merge(tbl.rows[3], 0, 1)
                tbl.rows[3].cells[0].merge(tbl.rows[5].cells[1]) # Vertical merge on left
                
                _set_labeled_cell(tbl.rows[3].cells[2], "Level:", plan.get("level", "6"))
                _merge(tbl.rows[3], 2, 3)
                _set_labeled_cell(tbl.rows[4].cells[2], "Class:", plan.get("class_code", ""))
                _merge(tbl.rows[4], 2, 3)
                _set_labeled_cell(tbl.rows[5].cells[2], "Number of Trainees:", plan.get("trainees_count", 20))
                _merge(tbl.rows[5], 2, 3)

                # Row 6, 7
                _set_labeled_cell(tbl.rows[6].cells[0], "Unit Code:", plan.get("unit_code", ""))
                _merge(tbl.rows[6], 0, 3)
                _set_labeled_cell(tbl.rows[7].cells[0], "Session Title:", plan.get("session_title", ""))
                _merge(tbl.rows[7], 0, 3)

                # Row 8
                _set_labeled_cell(tbl.rows[8].cells[0], "Learning outcome(s)")
                _merge(tbl.rows[8], 0, 1)
                
                outcomes_text = "\n".join([f"\u2610 {o}" for o in plan.get("learning_outcomes", [])])
                _set_labeled_cell(tbl.rows[8].cells[2], "By the end of the session the learner should be able to;\n", "\n" + outcomes_text)
                _merge(tbl.rows[8], 2, 3)

                # Row 9, 10
                resources_text = "\n".join([f"\u2610 {res}" for res in plan.get("resources", [])])
                _set_labeled_cell(tbl.rows[9].cells[0], "Resources (references, and learning aids)")
                _merge(tbl.rows[9], 0, 1)
                _merge(tbl.rows[9], 2, 3)
                tbl.rows[9].cells[2].text = "\n" + (resources_text or "\u2610 OS/Curriculum\n\u2610 Learning guides")
                
                _set_labeled_cell(tbl.rows[10].cells[0], "Safety requirements")
                _merge(tbl.rows[10], 0, 1)
                _merge(tbl.rows[10], 2, 3)
                tbl.rows[10].cells[2].text = "\n\u2610 " + plan.get("safety_requirements", "Adhere to lab safety rules.")

                # Row 11, 12, 13
                _set_header_cell(tbl.rows[11].cells[0], "Session presentation", align=1)
                _merge(tbl.rows[11], 0, 3)

                _set_header_cell(tbl.rows[12].cells[0], "1. Introduction", align=0)
                _merge(tbl.rows[12], 0, 3)
                tbl.rows[12].cells[0].paragraphs[0].add_run("\n" + plan.get("introduction", ""))

                _set_header_cell(tbl.rows[13].cells[0], "2. Session Delivery", align=0)
                _merge(tbl.rows[13], 0, 3)

                # Row 14 headers
                _set_header_cell(tbl.rows[14].cells[0], "Time (in minutes)", fill=None)
                _set_header_cell(tbl.rows[14].cells[1], "Trainer Activity", fill=None)
                _set_header_cell(tbl.rows[14].cells[2], "Learner Activity", fill=None)
                _set_header_cell(tbl.rows[14].cells[3], "Learning Check/Assessment", fill=None)

                # Steps
                for si, step in enumerate(steps):
                    rs = tbl.rows[15 + si]
                    rs.cells[0].text = str(step.get("time_minutes", ""))
                    rs.cells[1].text = step.get("trainer_activity", "")
                    rs.cells[2].text = step.get("learner_activity", "")
                    rs.cells[3].text = step.get("assessment", "")

                footer_start = 15 + n_step_rows

                # Footer rows
                _set_header_cell(tbl.rows[footer_start].cells[0], "3. Session review", fill=None, align=0)
                tbl.rows[footer_start].cells[1].text = plan.get("session_review", "Critique session: Highlight good use of alignment.")
                _merge(tbl.rows[footer_start], 1, 3)

                _set_header_cell(tbl.rows[footer_start+1].cells[0], "Assignment:", fill=None, align=0)
                tbl.rows[footer_start+1].cells[0].paragraphs[0].runs[0].italic = True
                tbl.rows[footer_start+1].cells[1].text = plan.get("assignment", "Sketch a logo idea on paper for a fictional \"Tech Company\" to be digitized next class.")
                tbl.rows[footer_start+1].cells[1].paragraphs[0].runs[0].italic = True
                _merge(tbl.rows[footer_start+1], 1, 3)

                total_time = plan.get("total_time", f"{sum(int(s.get('time_minutes', 0)) for s in steps)}Hrs")
                _set_header_cell(tbl.rows[footer_start+2].cells[0], f"TOTAL TIME:    {total_time}", fill=None, align=1)
                _merge(tbl.rows[footer_start+2], 0, 3)

                _set_header_cell(tbl.rows[footer_start+3].cells[0], "Session reflection", fill=None, align=0)
                _merge(tbl.rows[footer_start+3], 0, 3)

                # Blank row for reflection content
                tbl.rows[footer_start+4].cells[0].text = "\n" + plan.get("reflection", "") + "\n"
                _merge(tbl.rows[footer_start+4], 0, 3)

                _set_labeled_cell(tbl.rows[footer_start+5].cells[0], "Signature:", plan.get("trainer_name", "______________________________________________________"))
                _merge(tbl.rows[footer_start+5], 0, 3)

                f_stream = io.BytesIO()
                doc.save(f_stream)
                doc_base64 = base64.b64encode(f_stream.getvalue()).decode("utf-8")

                self.send_response(200)
                self.send_cors_headers()
                self.send_header("Content-Type", "application/json")
                self.end_headers()
                self.wfile.write(json.dumps({
                    "success": True,
                    "filename": f"MTTI_CUR05_SessionPlan_W{plan.get('week_number', 1)}.docx",
                    "file_data": doc_base64
                }).encode("utf-8"))

            elif path == "/api/export-learning-plan":
                # ── MTTI/F/CUR/02 Learning Plan ────────────────────────────────
                payload = json.loads(post_data.decode("utf-8"))
                lp = payload.get("payload", {})

                doc = docx.Document()
                section = doc.sections[0]
                section.top_margin    = docx.shared.Inches(0.6)
                section.bottom_margin = docx.shared.Inches(0.6)
                section.left_margin   = docx.shared.Inches(0.8)
                section.right_margin  = docx.shared.Inches(0.8)

                style = doc.styles['Normal']
                style.font.name = 'Maiandra GD'
                style.font.size = docx.shared.Pt(11)

                # Form code (top-left)
                pc = doc.add_paragraph()
                pc.alignment = 0
                pc.add_run("MTTI/F/CUR/02").bold = True

                # Title block
                current_dir  = os.path.dirname(os.path.abspath(__file__))
                project_root = os.path.dirname(current_dir)
                logo_path    = os.path.join(project_root, "client", "public", "mtti-logo.jpg")

                ph = doc.add_paragraph()
                ph.alignment = 1
                if os.path.exists(logo_path):
                    ph.add_run().add_picture(logo_path, width=docx.shared.Inches(1.0))
                rh1 = ph.add_run("\nMUKIRIA TECHNICAL TRAINING INSTITUTE")
                rh1.bold = True
                rh1.font.size = docx.shared.Pt(14)
                rh2 = ph.add_run("\nLEARNING PLAN TEMPLATE")
                rh2.bold = True
                rh2.underline = True
                rh2.font.size = docx.shared.Pt(12)

                # ── Metadata table (7 rows × 2 cols) ────────────────────────────
                meta = doc.add_table(rows=7, cols=2)
                meta.style = 'Table Grid'

                def _lp_row(row_idx, label_l, val_l, label_r="", val_r=""):
                    meta.rows[row_idx].cells[0].text = f"{label_l}{val_l}"
                    meta.rows[row_idx].cells[1].text = f"{label_r}{val_r}"
                    if meta.rows[row_idx].cells[0].paragraphs[0].runs:
                        meta.rows[row_idx].cells[0].paragraphs[0].runs[0].bold = True
                    if label_r and meta.rows[row_idx].cells[1].paragraphs[0].runs:
                        meta.rows[row_idx].cells[1].paragraphs[0].runs[0].bold = True

                _lp_row(0, "UNIT OF COMPETENCE: ", lp.get("unit_name", ""), "Unit Code: ", lp.get("unit_code", ""))
                _lp_row(1, "Name of Trainer: ", lp.get("trainer_name", "Alexander Kinoti"))
                _lp_row(2, "Institution: ", lp.get("institution", "Mukiria Technical Training Institute"), "Level: ", str(lp.get("level", "6")))
                _lp_row(3, "Date of Preparation: ", lp.get("date_prepared", ""), "Date of Revision: ", lp.get("date_revised", ""))
                _lp_row(4, "Number of Trainees: ", str(lp.get("trainees_count", 20)), "Class: ", lp.get("class_code", ""))
                _lp_row(5, "Duration: ", lp.get("duration", ""), "Series: ", lp.get("series", ""))
                _lp_row(6, "Qualification: ", lp.get("qualification", ""), "Department: ", lp.get("department", "Computing & Informatics"))

                # ── Weekly schedule table (9 cols) ───────────────────────────────
                weeks = lp.get("weeks", [])
                sched = doc.add_table(rows=len(weeks) + 1, cols=9)
                sched.style = 'Table Grid'

                col_headers = [
                    "Week", "Session No.", "Session Title", "Learning Outcome",
                    "Trainer\nActivities", "Trainee\nActivities",
                    "Resources & Refs", "Learning Checks/\nAssessments", "Reflections\n& Date"
                ]
                for ci, hdr in enumerate(col_headers):
                    cell = sched.rows[0].cells[ci]
                    cell.text = hdr
                    if cell.paragraphs[0].runs:
                        cell.paragraphs[0].runs[0].bold = True

                for wi, week_data in enumerate(weeks):
                    row = sched.rows[wi + 1]
                    w_num = week_data.get("week", wi + 1)
                    s_num = week_data.get("session_no", 1)
                    title = week_data.get("title", "")

                    if str(w_num) == "1" and int(s_num) == 1:
                        # Week 1 is for Admission and Orientation: merge the borders into one row
                        merged_cell = row.cells[0].merge(row.cells[8])
                        merged_cell.text = "WEEK 1: ADMISSION AND ORIENTATION"
                        p = merged_cell.paragraphs[0]
                        p.alignment = 1  # Center
                        if p.runs:
                            p.runs[0].bold = True
                            p.runs[0].font.name = 'Maiandra GD'
                            p.runs[0].font.size = docx.shared.Pt(11)
                        try:
                            from docx.oxml.ns import nsdecls
                            from docx.oxml import parse_xml
                            tcPr = merged_cell._tc.get_or_add_tcPr()
                            shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F2F2F2"/>')
                            tcPr.append(shd)
                        except Exception:
                            pass
                    else:
                        row.cells[0].text = str(w_num)
                        row.cells[1].text = str(s_num)
                        row.cells[2].text = title
                        row.cells[3].text = week_data.get("outcome", "")
                        row.cells[4].text = week_data.get("trainer_activities", "")
                        row.cells[5].text = week_data.get("trainee_activities", "")
                        row.cells[6].text = week_data.get("resources", "Refs:\n1. OS/Curr\n2. Learning Guide")
                        row.cells[7].text = week_data.get("assessments", "1. oral questioning\n2. written tests\n3. observation")
                        row.cells[8].text = week_data.get("reflections", "")

                # Signature line
                p_sig = doc.add_paragraph()
                p_sig.add_run(f"\nPrepared by: {lp.get('trainer_name', 'Alexander Kinoti')}     "
                               f"Date: {lp.get('date_prepared', '')}     Sign: _______________")

                f_stream = io.BytesIO()
                doc.save(f_stream)
                doc_base64 = base64.b64encode(f_stream.getvalue()).decode("utf-8")

                self.send_response(200)
                self.send_cors_headers()
                self.send_header("Content-Type", "application/json")
                self.end_headers()
                self.wfile.write(json.dumps({
                    "success": True,
                    "filename": f"MTTI_CUR02_LearningPlan_{lp.get('unit_code', '').replace('/', '_')}.docx",
                    "file_data": doc_base64
                }).encode("utf-8"))

            elif path == "/api/export-practical-exam":
                # ── CDACC Practical Assessment (4-table layout) ─────────────────
                payload = json.loads(post_data.decode("utf-8"))
                pe = payload.get("payload", {})

                doc = docx.Document()
                section = doc.sections[0]
                section.top_margin    = docx.shared.Inches(0.6)
                section.bottom_margin = docx.shared.Inches(0.6)
                section.left_margin   = docx.shared.Inches(0.8)
                section.right_margin  = docx.shared.Inches(0.8)

                style = doc.styles['Normal']
                style.font.name = 'Maiandra GD'
                style.font.size = docx.shared.Pt(11)

                current_dir  = os.path.dirname(os.path.abspath(__file__))
                project_root = os.path.dirname(current_dir)
                logo_path    = os.path.join(project_root, "client", "public", "mtti-logo.jpg")

                # Institution header
                ph = doc.add_paragraph()
                ph.alignment = 1
                if os.path.exists(logo_path):
                    ph.add_run().add_picture(logo_path, width=docx.shared.Inches(1.0))
                rh1 = ph.add_run("\nMUKIRIA TECHNICAL TRAINING INSTITUTE")
                rh1.bold = True
                rh1.font.size = docx.shared.Pt(14)
                dept_run = ph.add_run(f"\n{pe.get('department', 'DEPARTMENT OF COMPUTING AND INFORMATICS')}")
                dept_run.bold = True
                dept_run.font.size = docx.shared.Pt(11)
                pt1 = ph.add_run("\nINTERNAL EXAMINATION")
                pt1.bold = True
                pt2 = ph.add_run("\nPRACTICAL ASSESSMENT")
                pt2.bold = True
                pt2.underline = True

                if pe.get("time"):
                    pt = ph.add_run(f"\nTIME: {pe['time']}")
                    pt.bold = True

                # ── Table 0: Course info (6 rows × 2 cols) ──────────────────────
                t0 = doc.add_table(rows=6, cols=2)
                t0.style = 'Table Grid'
                info_rows = [
                    ("COURSE NAME:", pe.get("course_name", "")),
                    ("COURSE CODE:", pe.get("course_code", "")),
                    ("UNIT NAME",   pe.get("unit_name", "")),
                    ("UNIT CODE",   pe.get("unit_code", "")),
                    ("CLASS",       pe.get("class_code", "")),
                    ("SERIES",      pe.get("series", "")),
                ]
                for ri, (lbl, val) in enumerate(info_rows):
                    t0.rows[ri].cells[0].text = lbl
                    t0.rows[ri].cells[1].text = val
                    if t0.rows[ri].cells[0].paragraphs[0].runs:
                        t0.rows[ri].cells[0].paragraphs[0].runs[0].bold = True

                doc.add_paragraph()

                # CANDIDATE & ASSESSOR DETAILS header
                ph2 = doc.add_paragraph()
                ph2.alignment = 1
                ph2.add_run("CANDIDATE & ASSESSOR DETAILS").bold = True

                # ── Table 1: Candidate/Assessor (3 rows × 4 cols) ───────────────
                t1 = doc.add_table(rows=3, cols=4)
                t1.style = 'Table Grid'
                t1.rows[0].cells[0].text = "Candidate Name:"
                t1.rows[0].cells[2].text = "Reg. Code:"
                t1.rows[1].cells[0].text = "Assessor Name:"
                t1.rows[1].cells[2].text = "Date:"
                t1.rows[2].cells[0].text = "Venue:"
                for ri in range(3):
                    for ci in [0, 2]:
                        if t1.rows[ri].cells[ci].paragraphs[0].runs:
                            t1.rows[ri].cells[ci].paragraphs[0].runs[0].bold = True

                # Instructions
                doc.add_paragraph()
                pi = doc.add_paragraph()
                pi.add_run("INSTRUCTIONS:").bold = True
                for instr in pe.get("instructions", [
                    "The assessment shall include a practical session and an oral assessment.",
                    "The assessor will take photos and videos as you perform the tasks at critical points.",
                    "Ensure all required resources have been provided before beginning.",
                ]):
                    doc.add_paragraph(instr, style='List Bullet')

                # Project brief
                doc.add_paragraph()
                pb = doc.add_paragraph()
                pb.add_run("PROJECT BRIEF").bold = True
                doc.add_paragraph(pe.get("project_brief", "In this project, the candidate will be required to demonstrate competence in the unit."))

                # Elements covered
                doc.add_paragraph()
                pe_sec = doc.add_paragraph()
                pe_sec.add_run("SECTION 1\nElements Covered").bold = True
                for el in pe.get("elements", []):
                    doc.add_paragraph(el, style='List Number')

                # Tasks
                pt_hdr = doc.add_paragraph()
                pt_hdr.add_run("Tasks").bold = True
                for task in pe.get("tasks", []):
                    doc.add_paragraph(task, style='List Number')

                doc.add_paragraph()
                po = doc.add_paragraph()
                po.add_run("ORAL ASSESSMENT").bold = True
                doc.add_paragraph("The assessor will proceed to ask you oral questions based on the work done.")
                doc.add_page_break()

                # ── Table 2: Performance Checklist ──────────────────────────────
                criteria = pe.get("criteria", [])
                n_criteria = max(len(criteria), 10)
                t2 = doc.add_table(rows=n_criteria + 2, cols=5)
                t2.style = 'Table Grid'

                for ci, hdr in enumerate(["No.", "Performance Criterion", "Max Marks", "Awarded", "Remarks"]):
                    t2.rows[0].cells[ci].text = hdr
                    if t2.rows[0].cells[ci].paragraphs[0].runs:
                        t2.rows[0].cells[ci].paragraphs[0].runs[0].bold = True

                total_marks = 0
                for ri in range(n_criteria):
                    row = t2.rows[ri + 1]
                    row.cells[0].text = str(ri + 1)
                    if ri < len(criteria):
                        crit = criteria[ri]
                        row.cells[1].text = crit.get("criterion", "")
                        marks = crit.get("marks", 4)
                        row.cells[2].text = str(marks)
                        total_marks += marks
                    else:
                        row.cells[2].text = ""

                # Total row
                tr = t2.rows[n_criteria + 1]
                tr.cells[0].text = "Total score"
                tr.cells[0].merge(tr.cells[1])
                tr.cells[2].text = str(total_marks)
                if tr.cells[0].paragraphs[0].runs:
                    tr.cells[0].paragraphs[0].runs[0].bold = True

                doc.add_paragraph()

                # ── Table 3: Verdict / Feedback (5 rows × 4 cols) ───────────────
                t3 = doc.add_table(rows=5, cols=4)
                t3.style = 'Table Grid'

                v_cell = t3.rows[0].cells[0]
                v_cell.text = "The candidate was found to be\n\n        \u2610 Competent        \u2610 Not Yet Competent"
                v_cell.merge(t3.rows[0].cells[3])

                for ri, lbl in enumerate(["Feedback from candidate:", "Feedback to candidate:"], start=1):
                    fc = t3.rows[ri].cells[0]
                    fc.text = lbl
                    fc.merge(t3.rows[ri].cells[3])
                    if fc.paragraphs[0].runs:
                        fc.paragraphs[0].runs[0].bold = True

                t3.rows[3].cells[0].text = "Candidate's signature:"
                t3.rows[3].cells[2].text = "Date:"
                t3.rows[4].cells[0].text = "Assessor's signature:"
                t3.rows[4].cells[2].text = "Date:"
                for ri in [3, 4]:
                    for ci in [0, 2]:
                        if t3.rows[ri].cells[ci].paragraphs[0].runs:
                            t3.rows[ri].cells[ci].paragraphs[0].runs[0].bold = True

                f_stream = io.BytesIO()
                doc.save(f_stream)
                doc_base64 = base64.b64encode(f_stream.getvalue()).decode("utf-8")

                self.send_response(200)
                self.send_cors_headers()
                self.send_header("Content-Type", "application/json")
                self.end_headers()
                self.wfile.write(json.dumps({
                    "success": True,
                    "filename": f"MTTI_PracticalExam_{pe.get('unit_code', '').replace('/', '_')}.docx",
                    "file_data": doc_base64
                }).encode("utf-8"))

            elif path in ["/api/merge-row-pdf", "/api/export-row-book"]:
                payload = json.loads(post_data.decode("utf-8"))
                term = payload.get("term", "Term 3, 2026")
                pdf_bytes = generate_row_book(payload)
                pdf_base64 = base64.b64encode(pdf_bytes).decode("utf-8")
                safe_term = term.replace(" ", "_").replace(",", "").replace("/", "-")
                filename = f"MTTI_RoW_Book_{safe_term}.pdf"
                
                self.send_response(200)
                self.send_cors_headers()
                self.send_header("Content-Type", "application/json")
                self.end_headers()
                self.wfile.write(json.dumps({
                    "success": True,
                    "filename": filename,
                    "file_data": pdf_base64
                }).encode("utf-8"))

            elif path == "/api/generate-qr":
                payload = json.loads(post_data.decode("utf-8"))
                session_plan_id = payload.get("session_plan_id", "sp-1")
                target_url = payload.get("target_url")
                if not target_url:
                    target_url = f"http://localhost:3000/session/{session_plan_id}"
                payload["target_url"] = target_url
                
                qr_data_url = generate_qr_png_base64(target_url)
                poster_bytes = generate_door_poster(payload)
                poster_base64 = base64.b64encode(poster_bytes).decode("utf-8")
                filename = f"MTTI_Door_Notice_{session_plan_id}.pdf"
                
                self.send_response(200)
                self.send_cors_headers()
                self.send_header("Content-Type", "application/json")
                self.end_headers()
                self.wfile.write(json.dumps({
                    "success": True,
                    "session_plan_id": session_plan_id,
                    "target_url": target_url,
                    "qr_data_url": qr_data_url,
                    "poster_pdf": poster_base64,
                    "filename": filename
                }).encode("utf-8"))

            elif path == "/api/export-exam-results-docx":
                payload = json.loads(post_data.decode("utf-8"))
                doc_bytes = generate_exam_results_docx(payload)
                doc_base64 = base64.b64encode(doc_bytes).decode("utf-8")
                reg_num = str(payload.get("reg_number", "STUDENT")).replace("/", "_")
                unit_c = str(payload.get("unit_code", "EXAM")).replace("/", "_")
                filename = f"MTTI_Results_{reg_num}_{unit_c}.docx"

                self.send_response(200)
                self.send_cors_headers()
                self.send_header("Content-Type", "application/json")
                self.end_headers()
                self.wfile.write(json.dumps({
                    "success": True,
                    "status": "success",
                    "filename": filename,
                    "file_data": doc_base64,
                    "data": doc_base64
                }).encode("utf-8"))

            elif path == "/api/export-timetable-pdf":
                payload = json.loads(post_data.decode("utf-8"))
                pdf_bytes = generate_timetable_pdf(payload)
                pdf_base64 = base64.b64encode(pdf_bytes).decode("utf-8")
                title_slug = str(payload.get("schedule_title", "TIMETABLE")).replace(" ", "_").replace(":", "_").replace("/", "_")
                filename = f"MTTI_Timetable_{title_slug}.pdf"

                self.send_response(200)
                self.send_cors_headers()
                self.send_header("Content-Type", "application/json")
                self.end_headers()
                self.wfile.write(json.dumps({
                    "success": True,
                    "status": "success",
                    "filename": filename,
                    "file_data": pdf_base64,
                    "data": pdf_base64
                }).encode("utf-8"))

            elif path == "/api/export-exam-results-pdf":
                payload = json.loads(post_data.decode("utf-8"))
                pdf_bytes = generate_exam_results_pdf(payload)
                pdf_base64 = base64.b64encode(pdf_bytes).decode("utf-8")
                reg_num = str(payload.get("reg_number", "STUDENT")).replace("/", "_")
                unit_c = str(payload.get("unit_code", "EXAM")).replace("/", "_")
                filename = f"MTTI_Results_{reg_num}_{unit_c}.pdf"

                self.send_response(200)
                self.send_cors_headers()
                self.send_header("Content-Type", "application/json")
                self.end_headers()
                self.wfile.write(json.dumps({
                    "success": True,
                    "status": "success",
                    "filename": filename,
                    "file_data": pdf_base64,
                    "data": pdf_base64
                }).encode("utf-8"))

            else:
                self.send_response(404)
                self.send_cors_headers()
                self.end_headers()
                self.wfile.write(json.dumps({"error": "Not Found"}).encode("utf-8"))

        except Exception as e:
            import traceback
            traceback.print_exc()
            self.send_response(500)
            self.send_cors_headers()
            self.send_header("Content-Type", "application/json")
            self.end_headers()
            self.wfile.write(json.dumps({"error": str(e)}).encode("utf-8"))

def run(server_class=ThreadingHTTPServer, handler_class=APIHandler, port=8000):
    server_address = ("", port)
    httpd = server_class(server_address, handler_class)
    print(f"Starting MTTI Python standard HTTP server on port {port}...")
    httpd.serve_forever()

if __name__ == "__main__":
    run()
