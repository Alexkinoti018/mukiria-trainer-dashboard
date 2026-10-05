import docx
import os
import glob
import json

def inspect_ict6():
    path = r'D:\Curriculum and OS\Cycle 3\ICT6\6 Curriculum (1).docx'
    doc = docx.Document(path)
    print(f"Total tables: {len(doc.tables)}, Total paragraphs: {len(doc.paragraphs)}")
    
    for i in range(min(15, len(doc.tables))):
        t = doc.tables[i]
        header = [c.text.replace('\n', ' ').strip() for c in t.rows[0].cells]
        header_str = " | ".join(header).lower()
        if any(k in header_str for k in ['unit', 'code', 'title', 'duration', 'learning outcome']):
            print(f"\nTable {i} ({len(t.rows)} rows x {len(t.columns)} cols): {header}")
            for r in t.rows[1:10]:
                vals = [c.text.replace('\n', ' ').strip() for c in r.cells]
                print("   ", vals[:4])

if __name__ == '__main__':
    inspect_ict6()
