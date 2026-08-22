"""Debug the session plan export locally to find the 500 error."""
import sys
sys.path.insert(0, '.')
import os, io, base64, docx, docx.shared

plan = {
    'date': '14/07/2026', 'time_duration': '8AM-12PM', 'week_number': 1,
    'trainer_name': 'Alexander Kinoti', 'unit_name': 'Perform Graphic Design',
    'unit_code': 'ICT/OS/CS/CR/11/6/A', 'level': 6, 'class_code': 'ITECH6/M/24',
    'trainees_count': 12, 'session_title': 'Test Session',
    'learning_outcomes': ['Create shapes'], 'resources': ['Lab'],
    'delivery_steps': [
        {'time_minutes': '30', 'trainer_activity': 'Demo', 'learner_activity': 'Practice', 'assessment': 'Quiz'}
    ],
    'total_time': '1hr'
}

try:
    doc = docx.Document()

    from docx.oxml.ns import qn as _qn
    from docx.oxml import OxmlElement as _OxmlElement
    section = doc.sections[0]
    section.top_margin    = docx.shared.Inches(0.6)
    section.bottom_margin = docx.shared.Inches(0.6)
    section.left_margin   = docx.shared.Inches(0.8)
    section.right_margin  = docx.shared.Inches(0.8)

    style = doc.styles['Normal']
    style.font.name = 'Maiandra GD'
    style.font.size = docx.shared.Pt(11)

    logo_path = os.path.join('client', 'public', 'mtti-logo.jpg')

    p_code = doc.add_paragraph()
    p_code.alignment = 0
    p_code.add_run("MTTI/F/CUR/05").bold = True
    print("p_code OK")

    p_hdr = doc.add_paragraph()
    p_hdr.alignment = 1
    if os.path.exists(logo_path):
        p_hdr.add_run().add_picture(logo_path, width=docx.shared.Inches(1.0))
    r_inst = p_hdr.add_run("\nMUKIRIA TECHNICAL TRAINING INSTITUTE")
    r_inst.bold = True
    r_inst.font.size = docx.shared.Pt(14)
    print("p_hdr OK")

    steps = plan.get('delivery_steps', [])
    n_step_rows = max(len(steps), 1)
    total_rows = 8 + 1 + 1 + n_step_rows + 4 + 1
    print(f"Total rows: {total_rows}")
    tbl = doc.add_table(rows=total_rows, cols=5)
    tbl.style = 'Table Grid'
    print("Table created OK")

    def _bold_cell(cell, text):
        cell.text = ""
        run = cell.paragraphs[0].add_run(text)
        run.bold = True

    def _plain_cell(cell, text):
        cell.text = str(text)

    def _merge_right(row, start_col, end_col):
        row.cells[start_col].merge(row.cells[end_col])

    # Row 0
    _bold_cell(tbl.rows[0].cells[0], "Date:")
    _plain_cell(tbl.rows[0].cells[1], plan.get("date", ""))
    _bold_cell(tbl.rows[0].cells[2], "Time:")
    _plain_cell(tbl.rows[0].cells[3], plan.get("time_duration", ""))
    _merge_right(tbl.rows[0], 3, 4)
    print("Row 0 OK")

    # Row 1
    _bold_cell(tbl.rows[1].cells[0], "Week:")
    _plain_cell(tbl.rows[1].cells[1], str(plan.get("week_number", "")))
    _bold_cell(tbl.rows[1].cells[2], "Time:")
    _plain_cell(tbl.rows[1].cells[3], plan.get("time_duration", ""))
    _merge_right(tbl.rows[1], 3, 4)
    print("Row 1 OK")

    # Row 2
    _bold_cell(tbl.rows[2].cells[0], "Trainer name:")
    _plain_cell(tbl.rows[2].cells[1], plan.get("trainer_name", ""))
    _bold_cell(tbl.rows[2].cells[2], "Department:")
    _plain_cell(tbl.rows[2].cells[3], plan.get("department", "Computing & Informatics"))
    _merge_right(tbl.rows[2], 3, 4)
    print("Row 2 OK")

    # Rows 3-5
    for i in range(3):
        _bold_cell(tbl.rows[3+i].cells[0], "Unit of Competence:")
        _plain_cell(tbl.rows[3+i].cells[1], plan.get("unit_name", ""))
    right_meta = [
        ("Level:", str(plan.get("level", "6"))),
        ("Class:", plan.get("class_code", "")),
        ("Number of Trainees:", str(plan.get("trainees_count", 20))),
    ]
    for i, (lbl, val) in enumerate(right_meta):
        _bold_cell(tbl.rows[3+i].cells[2], lbl)
        _plain_cell(tbl.rows[3+i].cells[3], val)
        _merge_right(tbl.rows[3+i], 3, 4)
    print("Rows 3-5 OK")

    # Row 6
    _bold_cell(tbl.rows[6].cells[0], "Unit Code:")
    _plain_cell(tbl.rows[6].cells[1], plan.get("unit_code", ""))
    _merge_right(tbl.rows[6], 1, 4)
    print("Row 6 OK")

    # Row 7
    _bold_cell(tbl.rows[7].cells[0], "Session Title:")
    _plain_cell(tbl.rows[7].cells[1], plan.get("session_title", ""))
    _merge_right(tbl.rows[7], 1, 4)
    print("Row 7 OK")

    # Row 8: Outcomes
    r8 = tbl.rows[8]
    _bold_cell(r8.cells[0], "Learning outcome(s)")
    outcomes_text = "By the end of the session the learner should be able to;\n"
    outcomes_text += "\n".join([f"\u2022 {o}" for o in plan.get("learning_outcomes", [])])
    _plain_cell(r8.cells[1], outcomes_text)
    _merge_right(r8, 1, 4)
    print("Row 8 OK")

    # Row 9: Resources
    r9 = tbl.rows[9]
    _bold_cell(r9.cells[0], "Resources (references, and learning aids)")
    resources_text = "\n".join([f"\u2022 {res}" for res in plan.get("resources", [])])
    _plain_cell(r9.cells[1], resources_text or "\u2022 OS/Curriculum\n\u2022 Learning guides")
    _merge_right(r9, 1, 4)
    print("Row 9 OK")

    # Row 10: Safety
    r10 = tbl.rows[10]
    _bold_cell(r10.cells[0], "Safety requirements")
    _plain_cell(r10.cells[1], plan.get("safety_requirements", "\u2022 Adhere to lab safety rules."))
    _merge_right(r10, 1, 4)
    print("Row 10 OK")

    # Row 11: header
    r11 = tbl.rows[11]
    _bold_cell(r11.cells[0], "Session presentation")
    _merge_right(r11, 0, 4)
    print("Row 11 OK")

    # Row 12: Introduction
    r12 = tbl.rows[12]
    _bold_cell(r12.cells[0], "Introduction")
    _plain_cell(r12.cells[1], plan.get("introduction", ""))
    _merge_right(r12, 1, 4)
    print("Row 12 OK")

    # Row 13: Delivery header
    r13 = tbl.rows[13]
    _bold_cell(r13.cells[0], "2. Session Delivery")
    _merge_right(r13, 0, 4)
    print("Row 13 OK")

    # Row 14: Column headers
    r14 = tbl.rows[14]
    for ci, hdr in enumerate(["Time (in minutes)", "Trainer Activity", "Learner Activity", "Learning Check/Assessment", ""]):
        _bold_cell(r14.cells[ci], hdr)
    print("Row 14 OK")

    # Delivery steps
    for si, step in enumerate(steps):
        rs = tbl.rows[15 + si]
        _plain_cell(rs.cells[0], str(step.get("time_minutes", "")))
        _plain_cell(rs.cells[1], "")
        _plain_cell(rs.cells[2], step.get("trainer_activity", ""))
        _plain_cell(rs.cells[3], step.get("learner_activity", ""))
        _plain_cell(rs.cells[4], step.get("assessment", ""))
    print("Steps OK")

    footer_start = 15 + n_step_rows
    print(f"footer_start = {footer_start}, total_rows = {total_rows}")

    rf0 = tbl.rows[footer_start]
    _bold_cell(rf0.cells[0], "3. Session review")
    _plain_cell(rf0.cells[1], plan.get("session_review", ""))
    _merge_right(rf0, 1, 4)
    print("Footer 0 OK")

    rf1 = tbl.rows[footer_start + 1]
    _bold_cell(rf1.cells[0], "Assignment:")
    _plain_cell(rf1.cells[1], plan.get("assignment", ""))
    _merge_right(rf1, 1, 4)
    print("Footer 1 OK")

    rf2 = tbl.rows[footer_start + 2]
    _bold_cell(rf2.cells[0], "TOTAL TIME:")
    _plain_cell(rf2.cells[1], plan.get("total_time", ""))
    _merge_right(rf2, 1, 4)
    print("Footer 2 OK")

    rf3 = tbl.rows[footer_start + 3]
    _bold_cell(rf3.cells[0], "Session reflection")
    _plain_cell(rf3.cells[1], plan.get("reflection", ""))
    _merge_right(rf3, 1, 4)
    print("Footer 3 OK")

    rs_row = tbl.rows[footer_start + 4]
    _bold_cell(rs_row.cells[0], "Signature:")
    _plain_cell(rs_row.cells[1], plan.get("trainer_name", "Alexander Kinoti"))
    _merge_right(rs_row, 1, 2)
    _bold_cell(rs_row.cells[3], "Date:")
    _plain_cell(rs_row.cells[4], plan.get("signature_date", plan.get("date", "")))
    print("Signature row OK")

    f_stream = io.BytesIO()
    doc.save(f_stream)
    print(f"SAVE OK: {len(f_stream.getvalue())} bytes")

except Exception as e:
    import traceback
    traceback.print_exc()
