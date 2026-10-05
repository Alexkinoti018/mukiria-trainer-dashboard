import docx
import re
import json

def parse_curriculum_docx(path):
    doc = docx.Document(path)
    print(f"Loaded {path}, tables: {len(doc.tables)}, paragraphs: {len(doc.paragraphs)}")
    
    # Scan for table of units (Module summaries)
    units_list = []
    
    # 1. Look for tables with 'unit code', 'units title' or similar
    for t_idx, t in enumerate(doc.tables):
        header = [c.text.replace('\n', ' ').strip().lower() for c in t.rows[0].cells]
        hdr_str = " ".join(header)
        if ('unit code' in hdr_str or 'isced' in hdr_str or 'tvet cdacc' in hdr_str) and ('title' in hdr_str or 'name' in hdr_str):
            print(f"Found Units Table at index {t_idx} ({len(t.rows)} rows)")
            
            # Map column indices
            col_isced = -1
            col_cdacc = -1
            col_title = -1
            col_hours = -1
            
            for c_i, h in enumerate(header):
                if 'isced' in h or ('unit' in h and 'code' in h and 'cdacc' not in h):
                    if col_isced == -1: col_isced = c_i
                if 'cdacc' in h:
                    col_cdacc = c_i
                if 'title' in h or 'name' in h:
                    col_title = c_i
                if 'duration' in h or 'hour' in h:
                    col_hours = c_i
            
            if col_cdacc == -1 and col_isced != -1 and len(header) >= 3:
                # Often col 0 is isced, col 1 is cdacc, col 2 is title
                col_cdacc = 1
                col_title = 2
            
            print(f"Col mappings: isced={col_isced}, cdacc={col_cdacc}, title={col_title}, hours={col_hours}")
            for r_idx, r in enumerate(t.rows[1:]):
                cells = [c.text.replace('\n', ' ').strip() for c in r.cells]
                if not any(cells): continue
                # Skip header-like rows or sub-totals
                first_cell = cells[0].upper()
                if 'MODULE' in first_cell or 'TOTAL' in first_cell or 'SUB-TOTAL' in first_cell or 'COMMON' in first_cell or 'CORE' in first_cell:
                    continue
                
                isced_val = cells[col_isced] if col_isced >= 0 and col_isced < len(cells) else ""
                cdacc_val = cells[col_cdacc] if col_cdacc >= 0 and col_cdacc < len(cells) else ""
                title_val = cells[col_title] if col_title >= 0 and col_title < len(cells) else ""
                hours_val = cells[col_hours] if col_hours >= 0 and col_hours < len(cells) else ""
                
                if title_val and (isced_val or cdacc_val):
                    units_list.append({
                        "isced_code": isced_val,
                        "cdacc_code": cdacc_val,
                        "title": title_val,
                        "hours": hours_val,
                        "table_index": t_idx
                    })
    
    print(f"Extracted {len(units_list)} units from table listings:")
    for u in units_list:
        print(f"  [{u['cdacc_code'] or u['isced_code']}] {u['title']} ({u['hours']} hrs)")
        
    return units_list

if __name__ == '__main__':
    parse_curriculum_docx(r'D:\Curriculum and OS\Cycle 3\ICT6\6 Curriculum (1).docx')
