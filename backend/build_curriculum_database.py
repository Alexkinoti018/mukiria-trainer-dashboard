import docx
import os
import glob
import re
import json

CURRICULUM_DIR = r"D:\Curriculum and OS"

def clean_txt(text):
    if not text: return ""
    return re.sub(r'\s+', ' ', text).strip()

def clean_title(title):
    if not title: return ""
    t = title
    t = re.sub(r'ISCED\s+UNIT\s+CODE:?\s*[0-9A-Z\s/]+', '', t, flags=re.I)
    t = re.sub(r'TVET\s+CDACC\s+(?:UNIT\s+)?CODE:?\s*[0-9A-Z\s/\-]+', '', t, flags=re.I)
    t = re.sub(r'UNIT\s+CODE:?\s*[0-9A-Z\s/\-]+', '', t, flags=re.I)
    t = re.sub(r'DURATION\s+OF\s+UNIT:?\s*\d+\s*HOURS?', '', t, flags=re.I)
    t = re.sub(r'[\r\n\t]+', ' ', t)
    t = re.sub(r'^\d+\s*[-.]\s*', '', t)
    t = re.sub(r'\s+', ' ', t).strip(' :-\t')
    if not t or len(t) < 3: return ""
    if t.upper().startswith("MODULE") or t.upper().startswith("SUBTOTAL") or t.upper() in ["CORE", "COMMON", "BASIC"]:
        return ""
    # Capitalize nicely
    if t.isupper():
        t = t.title()
    return t

def parse_single_unit_docx(doc_path, dept, course, level):
    doc = docx.Document(doc_path)
    fname = os.path.basename(doc_path)
    print(f"Parsing Single Unit Doc: {fname}")
    
    # Read title & codes from paragraphs
    full_text = "\n".join([clean_txt(p.text) for p in doc.paragraphs if clean_txt(p.text)])
    
    # Unit Title
    title = ""
    for p in doc.paragraphs[:5]:
        t = clean_txt(p.text)
        if t and not any(k in t.upper() for k in ['ISCED', 'TVET', 'CDACC', 'DURATION', 'REPUBLIC', 'MINISTRY']):
            title = clean_title(t)
            if title: break
            
    if not title:
        title = os.path.splitext(fname)[0].replace('Curriculum For ', '').replace('OS For ', '').strip()
        title = clean_title(title)
        
    isced_m = re.search(r'ISCED\s+(?:UNIT\s+)?CODE:?\s*([0-9A-Z\s/]+)', full_text, re.I)
    cdacc_m = re.search(r'TVET\s+CDACC\s+(?:UNIT\s+)?CODE:?\s*([0-9A-Z\s/\-]+)', full_text, re.I)
    dur_m = re.search(r'Duration\s+(?:of\s+unit)?:?\s*(\d+)\s*hours', full_text, re.I)
    
    isced = clean_txt(isced_m.group(1)) if isced_m else ""
    cdacc = clean_txt(cdacc_m.group(1)) if cdacc_m else ""
    hrs = int(dur_m.group(1)) if dur_m else 120
    
    # Description
    desc_m = re.search(r'Unit\s+Description\s*\n*(.*?)(?=\n[A-Z\s]{4,}|\Z)', full_text, re.I | re.DOTALL)
    desc = clean_txt(desc_m.group(1)) if desc_m else f"This unit specifies competencies required to perform {title.lower()} according to TVET CDACC standards."
    
    # Outcomes from Table 0
    outcomes = []
    if doc.tables:
        t0 = doc.tables[0]
        for r in t0.rows[1:]:
            cells = [clean_txt(c.text) for c in r.cells]
            if not cells or not cells[0]: continue
            if 'total' in cells[0].lower(): continue
            h_val = 20
            if len(cells) > 1:
                d = re.findall(r'\d+', cells[1])
                if d: h_val = int(d[0])
            outcomes.append({
                "title": cells[0],
                "duration_hours": h_val,
                "content": [],
                "assessment_methods": ["Practical", "Written Test", "Observation"]
            })
            
    # Enrich content from Table 1, 2, 3...
    for t in doc.tables[1:]:
        if len(t.rows) < 1: continue
        hdr = [clean_txt(c.text).lower() for c in t.rows[0].cells]
        hdr_str = " ".join(hdr)
        if 'content' in hdr_str or 'suggested' in hdr_str:
            for r in t.rows[1:]:
                cells = [clean_txt(c.text) for c in r.cells]
                if len(cells) >= 2:
                    c_text = cells[1]
                    topics = [clean_txt(s) for s in re.split(r'[\n\r•;]+|\d+\.\d+\.?\s*', c_text) if len(clean_txt(s)) > 3]
                    if outcomes and topics:
                        outcomes[0]["content"].extend(topics[:6])
                        
    return [{
        "title": title,
        "isced_code": isced,
        "cdacc_code": cdacc,
        "duration_hours": hrs,
        "description": desc,
        "department": dept,
        "course": course,
        "level": level,
        "source_file": fname,
        "learning_outcomes": outcomes
    }]

def parse_docx_units(doc_path, dept, course, level):
    doc = docx.Document(doc_path)
    fname = os.path.basename(doc_path)
    print(f"Parsing Multi-Unit Doc: {fname} ({dept})")
    
    extracted_units = []
    
    for t_idx, t in enumerate(doc.tables):
        if len(t.rows) < 2: continue
        
        # Test row 0 and row 1 as header
        hdr_row_idx = 0
        header = [clean_txt(c.text).lower() for c in t.rows[0].cells]
        hdr_str = " ".join(header)
        
        if not (('unit' in hdr_str or 'isced' in hdr_str or 'cdacc' in hdr_str) and ('titl' in hdr_str or 'name' in hdr_str)):
            if len(t.rows) > 2:
                h1 = [clean_txt(c.text).lower() for c in t.rows[1].cells]
                h1_str = " ".join(h1)
                if (('unit' in h1_str or 'isced' in h1_str or 'cdacc' in h1_str) and ('titl' in h1_str or 'name' in h1_str)):
                    header = h1
                    hdr_row_idx = 1
                    hdr_str = h1_str
                    
        if ('unit' in hdr_str or 'isced' in hdr_str or 'cdacc' in hdr_str) and ('titl' in hdr_str or 'name' in hdr_str):
            col_isced = -1
            col_cdacc = -1
            col_title = -1
            col_hours = -1
            
            for c_i, h in enumerate(header):
                if 'isced' in h or ('unit' in h and 'code' in h and 'cdacc' not in h):
                    if col_isced == -1: col_isced = c_i
                if 'cdacc' in h:
                    col_cdacc = c_i
                if 'titl' in h or 'name' in h:
                    col_title = c_i
                if 'duration' in h or 'hour' in h or 'durati on' in h:
                    col_hours = c_i
                    
            if col_title != -1:
                for r in t.rows[hdr_row_idx + 1:]:
                    cells = [clean_txt(c.text) for c in r.cells]
                    if not any(cells): continue
                    first_cell = cells[0].upper()
                    if any(k in first_cell for k in ['MODULE', 'SUB-TOTAL', 'SUBTOTAL', 'TOTAL', 'MARKS', 'COMPETENCE', 'RATIO', 'CORE', 'COMMON', 'BASIC']):
                        continue
                            
                    raw_title = cells[col_title] if col_title < len(cells) else ""
                    title = clean_title(raw_title)
                    if not title or title.upper().startswith("MODULE") or title.upper() in ["CORE", "COMMON", "BASIC"]:
                        continue
                        
                    isced = cells[col_isced] if col_isced != -1 and col_isced < len(cells) else ""
                    cdacc = cells[col_cdacc] if col_cdacc != -1 and col_cdacc < len(cells) else ""
                    hrs_str = cells[col_hours] if col_hours != -1 and col_hours < len(cells) else "120"
                    
                    digits = re.findall(r'\d+', hrs_str)
                    hrs = int(digits[0]) if digits else 120
                    
                    if 'TOTAL' in isced.upper() or 'MODULE' in isced.upper(): isced = ""
                    if 'TOTAL' in cdacc.upper() or 'MODULE' in cdacc.upper(): cdacc = ""
                    
                    if title and (isced or cdacc or len(title) > 4):
                        if not any(u["title"].lower() == title.lower() for u in extracted_units):
                            extracted_units.append({
                                "title": title,
                                "isced_code": isced,
                                "cdacc_code": cdacc,
                                "duration_hours": hrs,
                                "department": dept,
                                "course": course,
                                "level": level,
                                "source_file": fname,
                                "learning_outcomes": []
                            })
                            
    print(f"  Extracted {len(extracted_units)} units from {fname}")
    return extracted_units

def build_complete_curriculum_catalog():
    print("=== Re-indexing D:\\Curriculum and OS for Complete Institutional Catalog ===")
    
    multi_unit_docs = [
        # ICT 6
        (r"D:\Curriculum and OS\Cycle 3\ICT6\6 Curriculum (1).docx", "Computing and Informatics", "Diploma in ICT (Level 6)", 6),
        # ICT 4
        (r"D:\Curriculum and OS\Cycle 3\ICT4\4 Curriculum.docx", "Computing and Informatics", "Artisan in ICT (Level 4)", 4),
        # Admin 5
        (r"D:\Curriculum and OS\Cycle 3\Admin 5\curr lv 5.docx", "Business and Office Administration", "Certificate in Office Administration (Level 5)", 5),
        # Civil Engineering 6
        (r"D:\Curriculum and OS\Cycle 3\CE\CIVIL ENGINEERING 6 MODULARIZED.docx", "Civil and Building Engineering", "Diploma in Civil Engineering (Level 6)", 6),
        # Land Survey 6
        (r"D:\Curriculum and OS\Cycle 3\LAND Survey\Land Survey Level 6\CURRICULUM.docx", "Land Survey and Geomatics", "Diploma in Land Survey (Level 6)", 6)
    ]
    
    single_unit_docs = [
        # Dedicated Computer Essentials Level 6
        (r"D:\Curriculum and OS\Cycle 3\OS & Curriculum For Computer Essential level 6\Curriculum For COMPUTER ESSENTIALS Level 6.docx", "Computing and Informatics", "Diploma in ICT (Level 6)", 6),
        # Dedicated Computer Operations Level 5
        (r"D:\Curriculum and OS\Cycle 3\OS & Curriculum for Computer Operation\Curriculum For COMPUTER OPERATIONS Level 5.docx", "Computing and Informatics", "Certificate in ICT (Level 5)", 5),
        # Dedicated Digital Literacy Level 6
        (r"D:\Curriculum and OS\Cycle 3\OS& CURRICULUM CYCLE 2 Demonstrate DIGITAL LITERACY\CURRICULUM FOR LEVEL 5 AND 6 DIGITAL LITERACY.docx", "Computing and Informatics", "Apply Digital Literacy (Common Unit)", 6)
    ]
    
    raw_units = []
    for path, dept, course, level in multi_unit_docs:
        if os.path.exists(path):
            raw_units.extend(parse_docx_units(path, dept, course, level))
            
    for path, dept, course, level in single_unit_docs:
        if os.path.exists(path):
            raw_units.extend(parse_single_unit_docx(path, dept, course, level))
            
    print(f"\nTotal raw units extracted: {len(raw_units)}")
    
    # Specific curated curriculum details for high-priority units taught by Trainer Kinoti
    curated_profiles = {
        "COMPUTER ESSENTIALS": {
            "outcomes": [
                {"title": "1. Manage computer devices", "duration_hours": 20, "content": ["Identify computer components and ports", "Booting procedures (cold/warm)", "Peripheral device connections (printer, scanner, projector)", "Safety and ergonomic setups"], "assessment_methods": ["Practical", "Oral questions", "Observation"]},
                {"title": "2. Manage desktop settings", "duration_hours": 30, "content": ["Operating system interface customization", "Display resolution and wallpaper configuration", "Accessibility features and keyboard shortcuts", "Control panel management"], "assessment_methods": ["Practical Assessment", "Product Checklist"]},
                {"title": "3. Perform file management", "duration_hours": 20, "content": ["Folder hierarchy creation and directory navigation", "File naming conventions and extensions", "Copying, moving, renaming, and deleting files", "Compression and archive utilities"], "assessment_methods": ["Practical", "Written Test"]},
                {"title": "4. Manage computer software", "duration_hours": 20, "content": ["Software installation and licensing", "Uninstallation and updates", "Antivirus configuration and system protection", "Default program configurations"], "assessment_methods": ["Practical Observation", "Portfolio of Evidence"]},
                {"title": "5. Perform online jobs", "duration_hours": 30, "content": ["Online freelancing platforms and account creation", "Data entry, transcription, and virtual assistance basics", "Safe digital payments and customer communication", "Job search and bidding techniques"], "assessment_methods": ["Practical Project", "Portfolio"]}
            ]
        },
        "COMPUTER OPERATIONS": {
            "outcomes": [
                {"title": "1. Process computerized word document", "duration_hours": 30, "content": ["Document creation and formatting", "Paragraph styling, headers, and footers", "Tables, graphics, and mail merge", "Page layout, proofing, and printing"], "assessment_methods": ["Practical Assessment", "Product Checklist"]},
                {"title": "2. Manipulate computerized spreadsheet", "duration_hours": 30, "content": ["Worksheet navigation and data entry", "Mathematical formulas and functions (SUM, AVERAGE, IF, VLOOKUP)", "Cell formatting and conditional formatting", "Charts, graphs, and data analysis"], "assessment_methods": ["Practical Assessment", "Simulations"]},
                {"title": "3. Maintain computerized database", "duration_hours": 30, "content": ["Database tables, fields, and data types", "Primary keys and relationships", "Queries (select, criteria, sorting)", "Forms and reporting"], "assessment_methods": ["Practical Project", "Written Test"]},
                {"title": "4. Prepare PowerPoint presentation", "duration_hours": 20, "content": ["Slide layouts and master slides", "Multimedia integration (audio, video, images)", "Transitions, animations, and timing", "Slide show delivery and audience engagement"], "assessment_methods": ["Oral Presentation", "Product Checklist"]},
                {"title": "5. Manipulate graphic application", "duration_hours": 25, "content": ["Image editing tools and layers", "Vector vs raster graphics", "Typography and visual design principles", "Exporting formats for web and print"], "assessment_methods": ["Practical Project", "Portfolio"]},
                {"title": "6. Perform online collaboration", "duration_hours": 15, "content": ["Cloud storage (Google Drive, OneDrive)", "Collaborative document editing", "Video conferencing etiquette and tools", "Calendar sharing and task scheduling"], "assessment_methods": ["Practical Assessment", "Observation"]}
            ]
        },
        "DIGITAL LITERACY": {
            "outcomes": [
                {"title": "1. Identify computer hardware components", "duration_hours": 10, "content": ["Input devices and output devices", "System unit and central processing unit (CPU)", "Storage media (HDD, SSD, Flash drives)", "Connecting peripherals securely"], "assessment_methods": ["Practical", "Oral Questioning"]},
                {"title": "2. Apply operating system functions", "duration_hours": 15, "content": ["Desktop navigation and window management", "File explorer and directory organization", "System settings and device management", "Basic troubleshooting and restart"], "assessment_methods": ["Practical Assessment", "Observation Checklist"]},
                {"title": "3. Apply word processing concepts", "duration_hours": 15, "content": ["Typing and formatting office memos and letters", "Spelling check and grammar review", "Tables, lists, and indentation", "Saving, exporting PDF, and printing"], "assessment_methods": ["Product Checklist", "Written Test"]},
                {"title": "4. Apply spreadsheet applications", "duration_hours": 10, "content": ["Entering numerical data and basic formulas", "Auto-sum, averages, and basic charts", "Sorting and filtering lists", "Printing worksheets cleanly"], "assessment_methods": ["Practical", "Product Checklist"]},
                {"title": "5. Utilize internet, email & cybersecurity", "duration_hours": 10, "content": ["Web browsing and search engines", "Sending and organizing professional emails with attachments", "Password safety, phishing awareness, and virus prevention", "Data privacy in technical workplace"], "assessment_methods": ["Practical Demonstration", "Written Test"]}
            ]
        },
        "NETWORK SETUP": {
            "outcomes": [
                {"title": "1. Terminate Computer network cables", "duration_hours": 50, "content": ["Network cabling standards (T568A / T568B)", "Cat5e/Cat6 cable stripping and crimping RJ-45", "Cable testing with continuity and wiremap testers", "Punch down on patch panels and keystone jacks"], "assessment_methods": ["Practical Observation", "Product Checklist"]},
                {"title": "2. Connect Computer network devices", "duration_hours": 70, "content": ["Network topology setup (Star, Bus, Mesh)", "Connecting switches, routers, and access points", "Structured cabling management and rack installation", "Safety guidelines in server rooms"], "assessment_methods": ["Practical", "Oral Questions"]},
                {"title": "3. Configure Network Devices", "duration_hours": 80, "content": ["IPv4 addressing and subnet masks", "DHCP configuration and static IP assignment", "Basic router and switch configuration", "Ping diagnostics, traceroute, and network verification"], "assessment_methods": ["Practical Project", "Simulations"]}
            ]
        },
        "NETWORK DESIGN": {
            "outcomes": [
                {"title": "1. Design computer network infrastructure", "duration_hours": 50, "content": ["Gather network requirements and user constraints", "Develop physical and logical network diagrams", "Subnetting and IP address planning (VLSM)", "Bandwidth estimation and equipment selection"], "assessment_methods": ["Practical Project", "Portfolio of Evidence"]},
                {"title": "2. Install and configure network infrastructure", "duration_hours": 60, "content": ["VLAN configuration and trunking (802.1Q)", "Inter-VLAN routing and routing protocols (OSPF, RIP)", "Wireless access point setup and SSID security", "Network operating system deployment"], "assessment_methods": ["Practical Assessment", "Simulations"]},
                {"title": "3. Manage and maintain computer network", "duration_hours": 50, "content": ["Network monitoring tools and traffic analysis", "Quality of Service (QoS) and access control lists (ACLs)", "Network backup and disaster recovery", "Performance optimization and firmware upgrades"], "assessment_methods": ["Practical Demonstration", "Written Test"]}
            ]
        },
        "WEBSITE": {
            "outcomes": [
                {"title": "1. Plan and design website structure", "duration_hours": 40, "content": ["User experience (UX) and wireframing", "Information architecture and site mapping", "Web accessibility (WCAG) and responsive principles", "Domain registration and hosting considerations"], "assessment_methods": ["Portfolio of Evidence", "Product Checklist"]},
                {"title": "2. Develop website frontend with HTML/CSS/JS", "duration_hours": 80, "content": ["Semantic HTML5 layout structure", "CSS3 styling, Flexbox, and Grid", "Responsive design with media queries", "Client-side scripting with JavaScript"], "assessment_methods": ["Practical Assessment", "Project"]},
                {"title": "3. Implement server-side and database integration", "duration_hours": 60, "content": ["Backend processing (REST APIs / PHP / Python / Node)", "Database connection and CRUD operations", "Form validation and user authentication", "State management and cookies/sessions"], "assessment_methods": ["Practical Assessment", "Simulations"]},
                {"title": "4. Test, deploy and maintain website", "duration_hours": 40, "content": ["Cross-browser testing and mobile responsiveness", "Web performance optimization and SEO", "FTP/Git deployment to live web server", "Content management and security updates"], "assessment_methods": ["Practical Project", "Observation Checklist"]}
            ]
        },
        "DATABASE": {
            "outcomes": [
                {"title": "1. Design relational database models", "duration_hours": 40, "content": ["Entity-Relationship (ER) modeling", "Database normalization (1NF, 2NF, 3NF)", "Primary and foreign key constraint definition", "Data integrity rules"], "assessment_methods": ["Written Test", "Portfolio of Evidence"]},
                {"title": "2. Create and populate database structures", "duration_hours": 50, "content": ["SQL DDL statements (CREATE, ALTER, DROP)", "SQL DML statements (INSERT, UPDATE, DELETE)", "Data type selection and index creation", "Importing data from external formats"], "assessment_methods": ["Practical Assessment", "Product Checklist"]},
                {"title": "3. Perform complex data querying and reporting", "duration_hours": 60, "content": ["SELECT statements with WHERE, GROUP BY, HAVING, ORDER BY", "Table JOINs (INNER, LEFT, RIGHT)", "Aggregate functions and subqueries", "Stored procedures and views"], "assessment_methods": ["Practical Assessment", "Written Test"]},
                {"title": "4. Administer, secure and backup database", "duration_hours": 50, "content": ["User accounts, permissions, and roles (GRANT/REVOKE)", "Database backup strategies (full, incremental)", "Restoration procedures and disaster recovery", "Performance tuning and query optimization"], "assessment_methods": ["Practical Demonstration", "Third Party Report"]}
            ]
        },
        "TOPOGRAPHICAL SURVEY": {
            "outcomes": [
                {"title": "1. Plan topographical survey works", "duration_hours": 30, "content": ["Site reconnaissance and flight/map review", "Selecting survey equipment (Total Station, GPS/GNSS, Levels)", "Establishing control network and benchmark referencing", "Safety considerations on site"], "assessment_methods": ["Oral Questioning", "Project Proposal"]},
                {"title": "2. Execute field measurements", "duration_hours": 60, "content": ["Total station setup, leveling, and centering", "Spot height and contour feature coding", "Traversing and error distribution calculations", "Field booking and electronic data logging"], "assessment_methods": ["Practical Fieldwork", "Observation Checklist"]},
                {"title": "3. Process survey data and produce topographical map", "duration_hours": 60, "content": ["Coordinate transformation and reduction", "CAD plotting and digital terrain modeling (DTM)", "Contour interpolation and spot leveling", "Survey report and cadastral plan preparation"], "assessment_methods": ["Product Checklist", "Portfolio of Evidence"]}
            ]
        },
        "OFFICE CORESPONDENCE": {
            "outcomes": [
                {"title": "1. Handle incoming mail", "duration_hours": 18, "content": ["Opening and inspecting incoming correspondence", "Sorting mail by department and urgency", "Entering records in mail log and remittances register", "Securing confidential correspondence"], "assessment_methods": ["Practical Assessment", "Observation"]},
                {"title": "2. Handle outgoing mail", "duration_hours": 27, "content": ["Categorizing outgoing mail (legal, confidential, general)", "Recording in dispatch register or mail delivery book", "Envelope addressing and postal guidelines", "Postage stamping and courier dispatch"], "assessment_methods": ["Practical Assessment", "Product Checklist"]},
                {"title": "3. Handle electronic correspondence", "duration_hours": 20, "content": ["Professional email drafting and netiquette", "Scanning hard-copy correspondence for e-distribution", "E-memo transmission and read receipts", "Archiving electronic communications"], "assessment_methods": ["Practical", "Written Test"]},
                {"title": "4. Maintain mail room equipment", "duration_hours": 15, "content": ["Franking machine and letter opening maintenance", "Paper shredder safety and servicing", "Documenting equipment malfunction", "Servicing records maintenance"], "assessment_methods": ["Observation Checklist", "Third Party Report"]}
            ]
        },
        "SITE SURVEY": {
            "outcomes": [
                {"title": "1. Prepare for site survey operations", "duration_hours": 30, "content": ["Review site drawings, benchmark datum, and survey briefs", "Inspect and calibrate optical levels, theodolites, and total stations", "Establish field survey safety and PPE protocols"], "assessment_methods": ["Practical", "Oral Questions"]},
                {"title": "2. Perform leveling and boundary setting-out", "duration_hours": 50, "content": ["Differential leveling and rise-and-fall booking", "Profile leveling, cross-sections, and cut/fill estimation", "Setting out building corners, gridlines, and batter boards"], "assessment_methods": ["Practical Fieldwork", "Observation Checklist"]},
                {"title": "3. Produce site survey plans and verification reports", "duration_hours": 40, "content": ["Plotting leveling data and contour profiles", "Verifying tolerance limits and checking closing errors", "Generating site handover survey reports"], "assessment_methods": ["Product Checklist", "Portfolio of Evidence"]}
            ]
        }
    }
    
    final_catalog = []
    seen = set()
    
    for u in raw_units:
        title = clean_title(u["title"])
        if not title: continue
        
        # Primary code
        code = u.get("cdacc_code") or u.get("isced_code")
        if not code:
            code = f"{title.upper()[:12].replace(' ', '_')}"
            
        key = f"{code}_{title.upper()}".replace(' ', '')
        if key in seen: continue
        seen.add(key)
        
        # Curated profile lookup
        matched_profile = None
        for k_p, prof in curated_profiles.items():
            if k_p in title.upper():
                matched_profile = prof
                break
                
        outcomes = u.get("learning_outcomes") or []
        if matched_profile:
            outcomes = matched_profile["outcomes"]
            
        if not outcomes:
            outcomes = [
                {
                    "title": f"1. Apply fundamental principles of {title}",
                    "duration_hours": max(20, u["duration_hours"] // 4),
                    "content": [f"Core concepts and terminology of {title}", "Safety, regulations, and standard procedures", "Equipment, materials, and workplace requirements"],
                    "assessment_methods": ["Practical", "Written Test", "Oral Questioning"]
                },
                {
                    "title": f"2. Execute practical operations in {title}",
                    "duration_hours": max(40, u["duration_hours"] // 2),
                    "content": [f"Setup, configuration, and preparation for {title}", f"Standard operational execution of {title} tasks", "Quality inspection and performance criteria verification"],
                    "assessment_methods": ["Practical Assessment", "Product Checklist", "Observation"]
                },
                {
                    "title": f"3. Perform maintenance, testing and troubleshooting",
                    "duration_hours": max(20, u["duration_hours"] // 4),
                    "content": ["Testing procedures and functional diagnostics", "Preventive maintenance and routine checks", "Documentation, logs, and reporting"],
                    "assessment_methods": ["Portfolio of Evidence", "Practical Demonstration"]
                }
            ]
            
        elements = [
            {
                "element_title": lo["title"],
                "performance_criteria": [
                    f"{idx+1}.1 Tools, equipment and materials are prepared according to task requirements.",
                    f"{idx+1}.2 Tasks are executed in compliance with CDACC occupational standards and safety procedures.",
                    f"{idx+1}.3 Quality checks and verification are documented accurately."
                ]
            } for idx, lo in enumerate(outcomes)
        ]
        
        # 10-week Learning Plan distribution
        weeks_breakdown = []
        for w in range(1, 11):
            if w == 1:
                w_title = "ADMISSION AND ORIENTATION"
                w_outcomes = ["Reporting, Admission and Orientation of trainees; familiarization with institute regulations, registration, and workshop safety."]
            elif w == 10:
                w_title = f"Practical Consolidation, Review & CDACC Assessment"
                w_outcomes = [f"Complete practical project portfolios and final competency assessment for {title}."]
            else:
                o_idx = min(len(outcomes) - 1, int((w - 2) / 8 * len(outcomes)))
                lo = outcomes[o_idx]
                w_title = lo["title"]
                w_outcomes = lo.get("content", [])[:3] or [f"Demonstrate competencies for {lo['title']}."]
                
            weeks_breakdown.append({
                "week": w,
                "title": w_title,
                "outcomes": w_outcomes,
                "resources": [
                    "Workshop workstations / lab equipment",
                    "Official CDACC Curriculum & Occupational Standards Guides",
                    "Training manuals & multimedia projector"
                ],
                "safety": "Strict adherence to workshop safety rules, ergonomic posture, and electrical safety precautions."
            })

        entry = {
            "id": code.replace('/', '-').replace(' ', '_'),
            "unit_code": code,
            "isced_code": u.get("isced_code", ""),
            "cdacc_code": u.get("cdacc_code", code),
            "unit_title": title,
            "department": u["department"],
            "course": u["course"],
            "level": u["level"],
            "duration_hours": u["duration_hours"],
            "description": u.get("description") or f"This unit specifies the competencies required to perform {title.lower()} in accordance with TVET CDACC occupational standards.",
            "source_file": u.get("source_file", ""),
            "learning_outcomes": outcomes,
            "elements": elements,
            "weeks_breakdown": weeks_breakdown,
            "suggested_resources": [
                "Computer workstations / workshop lab equipment",
                "Operating systems, tools and testing apparatus",
                "PPE gear and ergonomic furniture",
                "Official CDACC Curriculum Guides and Occupational Standards"
            ],
            "safety_protocols": "Strict adherence to workshop safety rules, electrical safety protocols, and anti-static precautions."
        }
        final_catalog.append(entry)

    print(f"\n==========================================")
    print(f"SUCCESS: Compiled {len(final_catalog)} official CDACC Units into Catalog!")
    print(f"==========================================")
    
    depts = {}
    for u in final_catalog:
        depts.setdefault(u["department"], []).append(u)
    for d, items in depts.items():
        print(f"\n[{d}] ({len(items)} units):")
        for u in items[:5]:
            print(f"   [{u['unit_code']}] {u['unit_title']} (L{u['level']} | {u['duration_hours']}h)")
        if len(items) > 5:
            print(f"   ... and {len(items)-5} more units")

    # Save to backend
    b_out = r"C:\Users\Administrator\Desktop\mukiria-trainer-dashboard\backend\data\curriculum_database.json"
    os.makedirs(os.path.dirname(b_out), exist_ok=True)
    with open(b_out, "w", encoding="utf-8") as f:
        json.dump(final_catalog, f, indent=2, ensure_ascii=False)
    print(f"\nSaved backend DB: {b_out}")

    # Save to client
    c_out = r"C:\Users\Administrator\Desktop\mukiria-trainer-dashboard\client\src\lib\curriculum_database.json"
    os.makedirs(os.path.dirname(c_out), exist_ok=True)
    with open(c_out, "w", encoding="utf-8") as f:
        json.dump(final_catalog, f, indent=2, ensure_ascii=False)
    print(f"Saved client DB: {c_out}")

if __name__ == '__main__':
    build_complete_curriculum_catalog()
